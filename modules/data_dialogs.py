import os, csv, json, re, tempfile, webbrowser
from datetime import datetime
from typing import Any, Dict, List, Optional
from PyQt5.QtCore import Qt, QTimer
from PyQt5.QtGui import QColor, QFont
from PyQt5.QtWidgets import (QApplication, QDialog, QWidget, QFrame,
    QVBoxLayout, QHBoxLayout, QFormLayout, QLabel, QLineEdit,
    QPushButton, QComboBox, QCheckBox, QTableWidget, QTableWidgetItem,
    QHeaderView, QAbstractItemView, QTreeWidget, QTreeWidgetItem,
    QTextBrowser, QPlainTextEdit, QDialogButtonBox, QMessageBox,
    QFileDialog, QMainWindow, QProgressBar, QGroupBox, QRadioButton)

from app_core.i18n import I18n
from app_core.config import RUNTIME_PATHS, AppConfig
from app_core.utils import JsonUtils
from services.database import DatabaseManager
from modules.print_engine import PrintEngine
from widgets.toast import ToastNotification


class ImportDialog(QDialog):
    def __init__(self, table: str = "", parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("import.title"))
        self.setMinimumSize(700, 550)
        self.resize(800, 600)
        self._source_columns: List[str] = []
        self._source_data: List[Dict[str, str]] = []
        self._target_table: str = table or "employees"
        self._build_ui()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(12)

        heading = QLabel(I18n._("import.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        file_layout = QHBoxLayout()
        self._file_path = QLineEdit()
        self._file_path.setPlaceholderText(I18n._("import.select"))
        file_layout.addWidget(self._file_path, 1)
        browse_btn = QPushButton("...")
        browse_btn.setFixedWidth(40)
        browse_btn.clicked.connect(self._browse_file)
        file_layout.addWidget(browse_btn)
        layout.addLayout(file_layout)

        table_layout = QHBoxLayout()
        table_layout.addWidget(QLabel(I18n._("common.table") + ":"))
        self._table_combo = QComboBox()
        self._populate_table_combo()
        idx = self._table_combo.findData(self._target_table)
        if idx >= 0:
            self._table_combo.setCurrentIndex(idx)
        table_layout.addWidget(self._table_combo)
        table_layout.addStretch()
        parse_btn = QPushButton(I18n._("common.preview"))
        parse_btn.clicked.connect(self._parse_file)
        table_layout.addWidget(parse_btn)
        layout.addLayout(table_layout)

        dup_layout = QHBoxLayout()
        dup_layout.addWidget(QLabel(I18n._("import.dup_strategy") + ":"))
        self._dup_strategy = QComboBox()
        self._dup_strategy.addItem(I18n._("import.dup_skip"), "skip")
        self._dup_strategy.addItem(I18n._("import.dup_update"), "update")
        self._dup_strategy.addItem(I18n._("import.dup_create"), "create")
        dup_layout.addWidget(self._dup_strategy)
        dup_layout.addStretch()
        layout.addLayout(dup_layout)

        self._preview_table = QTableWidget()
        self._preview_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._preview_table.setAlternatingRowColors(True)
        self._preview_table.verticalHeader().hide()
        layout.addWidget(self._preview_table, 1)

        map_heading = QLabel(I18n._("import.column_map"))
        map_heading.setStyleSheet("font-weight: 600; font-size: 14px;")
        layout.addWidget(map_heading)

        self._map_layout = QVBoxLayout()
        layout.addLayout(self._map_layout)

        self._mapping_widgets: List[Tuple[QLabel, QComboBox]] = []

        self._progress = QProgressBar()
        self._progress.setVisible(False)
        layout.addWidget(self._progress)

        btn_layout = QHBoxLayout()
        self._import_btn = QPushButton(I18n._("import.execute"))
        self._import_btn.setProperty("success", True)
        self._import_btn.clicked.connect(self._execute_import)
        self._import_btn.setEnabled(False)
        btn_layout.addWidget(self._import_btn)
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

    def _populate_table_combo(self) -> None:
        self._table_combo.clear()
        tables = [
            ("employees", I18n._("tab.employees")),
            ("violations", I18n._("tab.violations")),
            ("custom_ledger", I18n._("tab.custom_ledger")),
            ("incidents", I18n._("tab.incidents")),
            ("ppe", I18n._("tab.ppe")),
            ("training", I18n._("tab.training")),
            ("permits", I18n._("tab.permits")),
        ]
        for key, label in tables:
            self._table_combo.addItem(label, key)

    def _browse_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, I18n._("import.select"), "",
            "CSV (*.csv);;Excel (*.xlsx *.xls);;All files (*.*)")
        if path:
            self._file_path.setText(path)

    def _parse_file(self) -> None:
        path = self._file_path.text().strip()
        if not path or not os.path.isfile(path):
            QMessageBox.warning(self, I18n._("common.error"),
                                I18n._("error.file_not_found"))
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".csv":
                self._parse_csv(path)
            elif ext in (".xlsx", ".xls"):
                self._parse_excel(path)
            else:
                QMessageBox.warning(self, I18n._("common.error"),
                                    I18n._("error.import_failed"))
                return
            self._target_table = self._table_combo.currentData()
            self._build_mapping()
            self._import_btn.setEnabled(True)
        except Exception as e:
            QMessageBox.critical(self, I18n._("common.error"),
                                 f"{I18n._('error.import_failed')}: {e}")

    def _parse_csv(self, path: str) -> None:
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            self._source_columns = reader.fieldnames or []
            self._source_data = []
            for i, row in enumerate(reader):
                if i >= 100:
                    break
                self._source_data.append(row)
        self._show_preview()

    def _parse_excel(self, path: str) -> None:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows_data = list(ws.iter_rows(values_only=True))
            if not rows_data:
                return
            headers = [str(h) if h is not None else "" for h in rows_data[0]]
            self._source_columns = headers
            self._source_data = []
            for row in rows_data[1:101]:
                d = {}
                for i, h in enumerate(headers):
                    val = row[i] if i < len(row) else ""
                    d[h] = str(val) if val is not None else ""
                self._source_data.append(d)
            wb.close()
            self._show_preview()
        except ImportError:
            QMessageBox.warning(self, I18n._("common.error"),
                                "openpyxl " + I18n._("error.not_found"))

    def _show_preview(self) -> None:
        if not self._source_columns or not self._source_data:
            return
        self._preview_table.setColumnCount(len(self._source_columns))
        self._preview_table.setHorizontalHeaderLabels(self._source_columns)
        self._preview_table.setRowCount(min(len(self._source_data), 10))
        for i, row in enumerate(self._source_data[:10]):
            for j, col in enumerate(self._source_columns):
                self._preview_table.setItem(i, j,
                    QTableWidgetItem(row.get(col, "")))
        self._preview_table.resizeColumnsToContents()

    def _build_mapping(self) -> None:
        while self._map_layout.count():
            item = self._map_layout.takeAt(0)
            if item and item.widget():
                item.widget().deleteLater()
        self._mapping_widgets.clear()

        target_cols = self.db.get_columns_config(self._target_table)
        target_names = [c["name"] for c in target_cols if c["name"] != "ID"]

        for src_col in self._source_columns:
            row = QHBoxLayout()
            src_label = QLabel(f"  {src_col}:")
            src_label.setFixedWidth(160)
            row.addWidget(src_label)
            arrow = QLabel("→")
            arrow.setFixedWidth(20)
            arrow.setAlignment(Qt.AlignCenter)
            row.addWidget(arrow)
            combo = QComboBox()
            combo.addItem("— " + I18n._("import.skip") + " —", "")
            for tn in target_names:
                combo.addItem(tn, tn)
            guess = self._guess_mapping(src_col, target_names)
            if guess:
                idx = combo.findText(guess)
                if idx >= 0:
                    combo.setCurrentIndex(idx)
            row.addWidget(combo, 1)
            self._map_layout.addLayout(row)
            self._mapping_widgets.append((src_label, combo))

    def _guess_mapping(self, source: str, targets: List[str]) -> Optional[str]:
        sl = source.lower().strip()
        alias_map = {
            "фио": "ФИО", "fio": "ФИО", "full name": "ФИО", "name": "ФИО",
            "должность": "Должность", "position": "Должность",
            "фирма": "Фирма", "company": "Фирма", "организация": "Фирма",
            "телефон": "Телефон", "phone": "Телефон",
            "дата": "Дата", "date": "Дата",
            "статус": "Статус", "status": "Статус",
            "описание": "Описание", "description": "Описание",
            "штраф": "Штраф", "fine": "Штраф", "сумма": "Штраф",
            "категория": "Категория риска", "category": "Категория риска",
            "ответственный": "Ответственный", "responsible": "Ответственный",
            "подразделение": "Подразделение", "department": "Подразделение",
            "квалификация": "Квалификация", "qualification": "Квалификация",
        }
        for alias, target in alias_map.items():
            if alias in sl or sl in alias:
                if target in targets:
                    return target
        for t in targets:
            if t.lower().startswith(sl) or sl.startswith(t.lower()):
                return t
        return None

    def _execute_import(self) -> None:
        target_map: Dict[str, str] = {}
        for src_label, combo in self._mapping_widgets:
            src = src_label.text().strip().rstrip(":").strip()
            dst = combo.currentData()
            if dst:
                target_map[src] = dst
        if not target_map:
            QMessageBox.warning(self, I18n._("common.warning"),
                                I18n._("import.execute") + "?")
            return

        # Auto-create missing columns in target table
        target_cols = self.db.get_columns_config(self._target_table)
        existing_names = {c["name"] for c in target_cols}
        max_pos = max((c["position"] for c in target_cols), default=-1)
        for src, dst in target_map.items():
            if dst and dst not in existing_names:
                guessed_type = "Текст"
                lc = dst.lower()
                if any(x in lc for x in ["дата", "date", "срок", "deadline"]):
                    guessed_type = "Дата"
                elif any(x in lc for x in ["штраф", "fine", "сумма", "цена", "price"]):
                    guessed_type = "Число"
                elif any(x in lc for x in ["статус", "status"]):
                    guessed_type = "Статус"
                elif any(x in lc for x in ["фото", "photo", "media"]):
                    guessed_type = "Медиа"
                max_pos += 1
                self.db.execute(
                    "INSERT INTO columns_config (category, name, type, position) VALUES (?, ?, ?, ?)",
                    (self._target_table, dst, guessed_type, max_pos))
                self.db.conn.commit()
                existing_names.add(dst)

        dup_mode = self._dup_strategy.currentData()
        imported = 0
        updated = 0
        errors = 0
        import_log: List[str] = []
        total = len(self._source_data)
        self._progress.setVisible(True)
        self._progress.setMaximum(total)
        self._import_btn.setEnabled(False)
        QApplication.processEvents()

        for row_idx, row_data in enumerate(self._source_data):
            try:
                record: Dict[str, Any] = {}
                for src, dst in target_map.items():
                    record[dst] = row_data.get(src, "")
                record.setdefault("Фото", [])
                for c in self.db.get_columns_config(self._target_table):
                    if c["name"] not in record:
                        if c["type"] == "Статус":
                            record[c["name"]] = "Активно"
                        elif c["type"] in ("Дата", "Годен до", "Date", "Valid until"):
                            record[c["name"]] = ""
                        elif c["type"] == "Число":
                            record[c["name"]] = "0"
                        elif c["type"] == "Медиа":
                            record[c["name"]] = []
                        else:
                            record[c["name"]] = ""

                action = "created"
                target_id = 0
                if dup_mode != "create":
                    existing = self.db.find_duplicate(self._target_table, record)
                    if existing:
                        if dup_mode == "skip":
                            import_log.append(f"Строка {row_idx+1}: пропущен (дубликат ID={existing['id']})")
                            self._progress.setValue(row_idx + 1)
                            continue
                        merged = dict(existing.get("data_json", {}))
                        for k, v in record.items():
                            if isinstance(v, list):
                                if v:
                                    merged[k] = v
                            elif str(v).strip() != "":
                                merged[k] = v
                        self.db.save_json_record(self._target_table, existing["id"], merged)
                        updated += 1
                        action = "updated"
                        target_id = existing["id"]
                        import_log.append(f"Строка {row_idx+1}: обновлён ID={target_id} | {record.get('ФИО', record.get('Описание', record.get('Наименование', record.get('Наименование СИЗ', '?'))))[:60]}")
                        self._progress.setValue(row_idx + 1)
                        continue

                target_id = self.db.save_json_record(self._target_table, 0, record)
                imported += 1
                name_field = next((c["name"] for c in self._columns if c["name"] in ("ФИО", "Описание", "Наименование", "Наименование СИЗ", "Сотрудник")), "?")
                import_log.append(f"Строка {row_idx+1}: создан ID={target_id} | {record.get(name_field, '?')[:60]}")
            except Exception as e:
                errors += 1
                import_log.append(f"Строка {row_idx+1}: ОШИБКА — {e}")
            self._progress.setValue(row_idx + 1)
            QApplication.processEvents()

        self.db.log_event(f"Import: {imported} records, {updated} updated, {errors} errors",
                          "INFO", {"table": self._target_table})
        try:
            self.db.execute(
                "INSERT INTO import_history (table_name, source_file, imported, updated, errors, details) VALUES (?, ?, ?, ?, ?, ?)",
                (self._target_table, getattr(self, "_source_file", ""), imported, updated, errors, "\n".join(import_log[:50])))
            self.db.conn.commit()
        except Exception:
            pass
        self._show_import_log_dialog(imported, updated, errors, import_log)

    def _show_import_log_dialog(self, imported: int, updated: int, errors: int, log: List[str]) -> None:
        dlg = QDialog(self)
        dlg.setWindowTitle(I18n._("import.log_title"))
        dlg.resize(700, 450)
        layout = QVBoxLayout(dlg)
        summary = QLabel(
            f"<b>{I18n._('import.success').format(count=imported)}</b>"
            + (f" | {I18n._('import.updated')}: {updated}" if updated else "")
            + (f" | {I18n._('error.generic')}: {errors}" if errors else "")
        )
        summary.setTextFormat(Qt.RichText)
        summary.setWordWrap(True)
        layout.addWidget(summary)
        text = QPlainTextEdit()
        text.setReadOnly(True)
        text.setPlainText("\n".join(log))
        text.setLineWrapMode(QPlainTextEdit.NoWrap)
        layout.addWidget(text, 1)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Save)
        btns.accepted.connect(dlg.accept)
        btns.button(QDialogButtonBox.Save).clicked.connect(lambda: self._save_import_log(log))
        layout.addWidget(btns)
        dlg.exec_()

    def _save_import_log(self, log: List[str]) -> None:
        path, _ = QFileDialog.getSaveFileName(self, I18n._("import.save_log"),
                                              "import_log.txt", "Text (*.txt)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                f.write("\n".join(log))
            ToastNotification.notify(I18n._("export.success").format(path=path), "success", 3000)


class ExportDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("export.title"))
        self.setMinimumSize(400, 250)
        self.resize(450, 280)
        self._build_ui()

    def _populate_export_tables(self) -> None:
        self._table_combo.clear()
        tables = [
            ("employees", I18n._("tab.employees")),
            ("violations", I18n._("tab.violations")),
            ("custom_ledger", I18n._("tab.custom_ledger")),
            ("incidents", I18n._("tab.incidents")),
            ("ppe", I18n._("tab.ppe")),
            ("training", I18n._("tab.training")),
            ("permits", I18n._("tab.permits")),
        ]
        for key, label in tables:
            self._table_combo.addItem(label, key)

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 20, 24, 20)
        layout.setSpacing(14)

        heading = QLabel(I18n._("export.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        form = QFormLayout()
        form.setSpacing(10)

        self._table_combo = QComboBox()
        self._populate_export_tables()
        form.addRow(I18n._("common.table") + ":", self._table_combo)

        self._format_combo = QComboBox()
        self._format_combo.addItem("CSV (.csv)", "csv")
        self._format_combo.addItem("Excel (.xlsx)", "xlsx")
        form.addRow(I18n._("common.format") + ":", self._format_combo)

        layout.addLayout(form)

        btn_layout = QHBoxLayout()
        self._export_btn = QPushButton(I18n._("common.export"))
        self._export_btn.setProperty("success", True)
        self._export_btn.clicked.connect(self._do_export)
        btn_layout.addWidget(self._export_btn)
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

        self._status_label = QLabel()
        self._status_label.setStyleSheet("font-size: 12px;")
        self._status_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self._status_label)

    def _do_export(self) -> None:
        table = self._table_combo.currentData()
        fmt = self._format_combo.currentData()
        ext = ".csv" if fmt == "csv" else ".xlsx"
        default_name = f"{table}_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        default_path = os.path.join(RUNTIME_PATHS.export_dir, default_name)
        path, _ = QFileDialog.getSaveFileName(
            self, I18n._("common.export"), default_path,
            "CSV (*.csv)" if fmt == "csv" else "Excel (*.xlsx)")
        if not path:
            return

        try:
            if fmt == "csv":
                result = self.db.export_to_csv(table, path)
            else:
                result = self._export_excel(table, path)
            self._status_label.setText(
                I18n._("export.success").format(path=os.path.basename(result)))
            ToastNotification.notify(I18n._("export.success").format(
                path=os.path.basename(result)), "success", 5000)
        except Exception as e:
            self._status_label.setText(I18n._("export.error").format(error=str(e)))
            ToastNotification.notify(I18n._("export.error").format(error=str(e)),
                                   "error", 5000)

    def _export_excel(self, table: str, path: str) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            QMessageBox.warning(self, I18n._("common.error"),
                                "openpyxl " + I18n._("error.not_found"))
            return path

        records = self.db.get_json_records(table)
        columns = self.db.get_columns_config(table)
        headers = [c["name"] for c in columns]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = table

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1A237E", end_color="1A237E",
                                  fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"))

        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        for ri, rec in enumerate(records, 2):
            dj = rec.get("data_json", {})
            for ci, col in enumerate(columns, 1):
                name = col["name"]
                val = dj.get(name, "")
                if isinstance(val, list):
                    val = ", ".join(val) if val else ""
                cell = ws.cell(row=ri, column=ci, value=str(val))
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        wb.save(path)
        self.db.log_event("Excel export completed", "INFO",
                          {"table": table, "path": path})
        return path


class GlobalSearchDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("search.global"))
        self.setMinimumSize(700, 500)
        self.resize(800, 550)
        self._build_ui()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(12)

        heading = QLabel(I18n._("search.global"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        search_layout = QHBoxLayout()
        self._search_edit = QLineEdit()
        self._search_edit.setProperty("search", True)
        self._search_edit.setPlaceholderText(I18n._("search.global_placeholder"))
        self._search_edit.textChanged.connect(self._do_search)
        search_layout.addWidget(self._search_edit, 1)
        self._regex_cb = QCheckBox(I18n._("search.regex"))
        self._regex_cb.toggled.connect(self._do_search)
        search_layout.addWidget(self._regex_cb)
        layout.addLayout(search_layout)

        self._results_tree = QTreeWidget()
        self._results_tree.setHeaderLabels([I18n._("search.global"),
                                            I18n._("common.description")])
        self._results_tree.setAlternatingRowColors(True)
        self._results_tree.setColumnWidth(0, 200)
        self._results_tree.itemDoubleClicked.connect(self._on_result_clicked)
        layout.addWidget(self._results_tree)

        self._info_label = QLabel()
        self._info_label.setStyleSheet("font-size: 12px;")
        layout.addWidget(self._info_label)

        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        layout.addWidget(close_btn, 0, Qt.AlignCenter)

    def _do_search(self) -> None:
        text = self._search_edit.text().strip()
        self._results_tree.clear()
        if not text:
            self._info_label.setText("")
            return

        use_regex = self._regex_cb.isChecked()
        pattern = None
        if use_regex:
            try:
                pattern = re.compile(text, re.IGNORECASE)
            except re.error:
                self._info_label.setText(I18n._("error.generic"))
                return

        total = 0
        tables = [
            ("employees", I18n._("tab.employees"), "ФИО"),
            ("violations", I18n._("tab.violations"), "Описание"),
            ("custom_ledger", I18n._("tab.custom_ledger"), "Описание"),
            ("incidents", I18n._("tab.incidents"), "Описание"),
            ("ppe", I18n._("tab.ppe"), "Наименование СИЗ"),
            ("training", I18n._("tab.training"), "Наименование"),
            ("permits", I18n._("tab.permits"), "Описание работ"),
            ("companies", I18n._("tab.companies"), "name"),
        ]
        for table, label, title_field in tables:
            records = []
            if table == "companies":
                records_data = self.db.fetch_all(
                    "SELECT id, name, address, contact, data_json FROM companies")
                for r in records_data:
                    rec = dict(r)
                    rec["data_json"] = JsonUtils.loads(r.get("data_json", "{}"))
                    records.append(rec)
            else:
                records = self.db.get_json_records(table)

            matched = []
            for rec in records:
                dj = rec.get("data_json", {})
                searchable = {**dj}
                if table == "companies":
                    searchable["name"] = rec.get("name", "")
                    searchable["address"] = rec.get("address", "")
                found = False
                for val in searchable.values():
                    sval = str(val)
                    if use_regex:
                        if pattern and pattern.search(sval):
                            found = True
                            break
                    else:
                        if text.lower() in sval.lower():
                            found = True
                            break
                if found:
                    matched.append(rec)

            if matched:
                parent_item = QTreeWidgetItem(self._results_tree)
                parent_item.setText(0, f"{label} ({len(matched)})")
                parent_item.setText(1, "")
                parent_item.setExpanded(False)
                for rec in matched:
                    dj = rec.get("data_json", {})
                    title_val = dj.get(title_field, dj.get("name", ""))
                    desc = dj.get("Описание", dj.get("address", ""))
                    if isinstance(desc, str) and len(desc) > 80:
                        desc = desc[:80] + "..."
                    child = QTreeWidgetItem(parent_item)
                    child.setText(0, str(title_val)[:60])
                    child.setText(1, str(desc)[:120])
                    child.setData(0, Qt.UserRole, table)
                    child.setData(0, Qt.UserRole + 1, rec.get("id", 0))
                total += len(matched)

        self._info_label.setText(
            f"{I18n._('common.filter')}: {total}")

    def _on_result_clicked(self, item: QTreeWidgetItem, col: int) -> None:
        table = item.data(0, Qt.UserRole)
        rid = item.data(0, Qt.UserRole + 1)
        if table and rid and isinstance(self.parent(), QMainWindow):
            mw = self.parent()
            tab_keys = {"employees": "tab.employees", "violations": "tab.violations",
                        "companies": "tab.companies", "custom_ledger": "tab.custom_ledger",
                        "incidents": "tab.incidents", "ppe": "tab.ppe",
                        "training": "tab.training", "permits": "tab.permits"}
            key = tab_keys.get(table)
            if key and hasattr(mw, '_tabs_data') and key in mw._tabs_data:
                idx, tab_w = mw._tabs_data[key]
                mw._tab_widget.setCurrentIndex(idx)
                if hasattr(tab_w, '_load_data'):
                    tab_w._load_data()
                if hasattr(tab_w, 'focus_record'):
                    tab_w.focus_record(rid)
            ToastNotification.notify(f"{table} #{rid}", "info", 3000)


class ReportDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("report.title"))
        self.setMinimumSize(500, 450)
        self.resize(550, 500)
        self._build_ui()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 20, 24, 20)
        layout.setSpacing(14)

        heading = QLabel(I18n._("report.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        form = QFormLayout()
        form.setSpacing(10)

        self._company_combo = QComboBox()
        self._company_combo.addItem(I18n._("report.all_companies"), "")
        for c in self.db.get_companies():
            self._company_combo.addItem(c["name"], c["name"])
        form.addRow(I18n._("report.company") + ":", self._company_combo)

        self._include_emp = QCheckBox(I18n._("report.employees"))
        self._include_emp.setChecked(True)
        form.addRow("", self._include_emp)

        self._include_viol = QCheckBox(I18n._("report.violations"))
        self._include_viol.setChecked(True)
        form.addRow("", self._include_viol)

        self._include_fines = QCheckBox(I18n._("report.fines"))
        self._include_fines.setChecked(True)
        form.addRow("", self._include_fines)

        layout.addLayout(form)

        btn_layout = QHBoxLayout()
        self._preview_btn = QPushButton(I18n._("common.preview"))
        self._preview_btn.clicked.connect(self._generate_report)
        btn_layout.addWidget(self._preview_btn)
        self._export_html_btn = QPushButton(I18n._("export.title") + " HTML")
        self._export_html_btn.clicked.connect(self._export_html)
        btn_layout.addWidget(self._export_html_btn)
        self._export_excel_btn = QPushButton(I18n._("export.title") + " Excel")
        self._export_excel_btn.clicked.connect(self._export_excel_report)
        btn_layout.addWidget(self._export_excel_btn)
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

        self._status_lbl = QLabel()
        self._status_lbl.setStyleSheet("font-size: 12px;")
        self._status_lbl.setWordWrap(True)
        layout.addWidget(self._status_lbl)

        self._last_html: str = ""

    def _generate_report(self) -> str:
        company = self._company_combo.currentData() or ""
        include_emp = self._include_emp.isChecked()
        include_viol = self._include_viol.isChecked()
        include_fines = self._include_fines.isChecked()
        html = PrintEngine.render_report(company, include_emp, include_viol, include_fines)
        self._last_html = html
        preview_path = os.path.join(tempfile.gettempdir(), "suot_report_preview.html")
        with open(preview_path, "w", encoding="utf-8") as f:
            f.write(html)
        webbrowser.open(f"file://{preview_path}")
        self._status_lbl.setText(I18n._("report.generated"))
        return html

    def _export_html(self) -> None:
        if not self._last_html:
            self._generate_report()
        path, _ = QFileDialog.getSaveFileName(
            self, I18n._("common.export"),
            os.path.join(RUNTIME_PATHS.export_dir,
                         f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"),
            "HTML (*.html)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                f.write(self._last_html)
            self._status_lbl.setText(I18n._("export.success").format(path=os.path.basename(path)))

    def _export_excel_report(self) -> None:
        if not self._last_html:
            self._generate_report()
        path, _ = QFileDialog.getSaveFileName(
            self, I18n._("common.export"),
            os.path.join(RUNTIME_PATHS.export_dir,
                         f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"),
            "Excel (*.xlsx)")
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            QMessageBox.warning(self, I18n._("common.error"), "openpyxl required")
            return

        company = self._company_combo.currentData() or ""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = I18n._("report.title")

        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="1A237E", end_color="1A237E",
                                  fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"))

        # Title
        ws.cell(row=1, column=1,
                value=f"{I18n._('report.title')} — {company or I18n._('report.all_companies')}")
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)
        ws.merge_cells("A1:E1")

        now = datetime.now()
        ws.cell(row=2, column=1, value=now.strftime("%d.%m.%Y %H:%M"))
        ws.merge_cells("A2:E2")

        # Summary header
        row = 4
        headers = [I18n._("company.name"), I18n._("company.employees_count"),
                   I18n._("company.violations_count"), I18n._("company.fines_total"),
                   I18n._("stat.overdue_total")]
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        if company:
            companies_data = [{"name": company}]
        else:
            companies_data = self.db.get_companies()

        for ri, comp in enumerate(companies_data, row + 1):
            cname = comp.get("name", "")
            emp_c = sum(1 for e in self.db.get_json_records("employees")
                        if e.get("data_json", {}).get("Фирма") == cname)
            viol_c = 0
            fines = 0.0
            overdue = 0
            for v in self.db.get_json_records("violations"):
                dj = v.get("data_json", {})
                if dj.get("Фирма") == cname:
                    viol_c += 1
                    try:
                        fines += float(str(dj.get("Штраф", "0"))
                                       .replace(" ", "").replace(",", "."))
                    except Exception:
                        pass
                    dl = dj.get("Срок устранения", "")
                    try:
                        p = dl.split(".")
                        if len(p) == 3:
                            if datetime(int(p[2]), int(p[1]), int(p[0])) < now:
                                overdue += 1
                    except Exception:
                        pass
            vals = [cname, emp_c, viol_c, fines, overdue]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="right" if ci > 1 else "left")

        ws.column_dimensions["A"].width = 25
        for c in "BCDE":
            ws.column_dimensions[c].width = 18
        wb.save(path)
        self._status_lbl.setText(I18n._("export.success").format(path=os.path.basename(path)))
        ToastNotification.notify(I18n._("common.success"), "success", 3000)


class QuickReportDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("report.title"))
        self.setMinimumSize(900, 600)
        self.resize(1000, 650)
        self._build_ui()
        self._generate()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(12)

        heading = QLabel(I18n._("report.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        self._browser = QTextBrowser()
        self._browser.setOpenExternalLinks(True)
        layout.addWidget(self._browser)

        btn_layout = QHBoxLayout()
        export_btn = QPushButton(I18n._("common.export") + " HTML")
        export_btn.clicked.connect(self._export_html)
        btn_layout.addWidget(export_btn)
        export_excel_btn = QPushButton(I18n._("common.export") + " Excel")
        export_excel_btn.clicked.connect(self._export_excel)
        btn_layout.addWidget(export_excel_btn)
        print_btn = QPushButton(I18n._("common.print"))
        print_btn.clicked.connect(self._print)
        btn_layout.addWidget(print_btn)
        refresh_btn = QPushButton(I18n._("common.refresh"))
        refresh_btn.clicked.connect(self._generate)
        btn_layout.addWidget(refresh_btn)
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

        self._html: str = ""

    def _generate(self) -> None:
        html = PrintEngine.render_report("", True, True, True)
        self._html = html
        self._browser.setHtml(html)

    def _export_html(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self, I18n._("common.export"),
            os.path.join(RUNTIME_PATHS.export_dir,
                         f"quick_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"),
            "HTML (*.html)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                f.write(self._html)
            ToastNotification.notify(I18n._("common.success"), "success", 3000)

    def _export_excel(self) -> None:
        try:
            import openpyxl
        except ImportError:
            QMessageBox.warning(self, I18n._("common.error"), "openpyxl required")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, I18n._("common.export"),
            os.path.join(RUNTIME_PATHS.export_dir,
                         f"quick_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"),
            "Excel (*.xlsx)")
        if path:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Report"
            now = datetime.now()
            ws.cell(row=1, column=1, value=f"{I18n._('report.title')} — {now.strftime('%d.%m.%Y')}")
            ws.merge_cells("A1:E1")
            ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True, size=14)

            headers = [I18n._("company.name"), I18n._("company.employees_count"),
                       I18n._("company.violations_count"), I18n._("company.fines_total"),
                       I18n._("stat.overdue_total")]
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(row=3, column=ci, value=h)
                cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
                cell.fill = openpyxl.styles.PatternFill(start_color="1A237E",
                                                          end_color="1A237E",
                                                          fill_type="solid")
            row = 4
            for comp in self.db.get_companies():
                cname = comp.get("name", "")
                data = {"name": cname}
                ws.cell(row=row, column=1, value=cname)
                ws.cell(row=row, column=2,
                        value=sum(1 for e in self.db.get_json_records("employees")
                                  if e.get("data_json", {}).get("Фирма") == cname))
                viols = [v for v in self.db.get_json_records("violations")
                         if v.get("data_json", {}).get("Фирма") == cname]
                ws.cell(row=row, column=3, value=len(viols))
                fines = 0.0
                overdue = 0
                for v in viols:
                    dj = v.get("data_json", {})
                    try:
                        fines += float(str(dj.get("Штраф", "0"))
                                       .replace(" ", "").replace(",", "."))
                    except Exception:
                        pass
                    dl = dj.get("Срок устранения", "")
                    try:
                        p = dl.split(".")
                        if len(p) == 3 and datetime(int(p[2]), int(p[1]), int(p[0])) < now:
                            overdue += 1
                    except Exception:
                        pass
                ws.cell(row=row, column=4, value=fines)
                ws.cell(row=row, column=5, value=overdue)
                row += 1

            ws.column_dimensions["A"].width = 25
            for c in "BCDE":
                ws.column_dimensions[c].width = 18
            wb.save(path)
            ToastNotification.notify(I18n._("common.success"), "success", 3000)

    def _print(self) -> None:
        PrintEngine.print_document(self._html)

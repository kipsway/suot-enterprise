import csv
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from PyQt5.QtCore import Qt, QPoint, QTimer, QObject
from PyQt5.QtGui import QColor
from PyQt5.QtWidgets import (QApplication, QWidget, QDialog, QVBoxLayout,
                             QHBoxLayout, QLabel, QPushButton, QLineEdit,
                             QCheckBox, QComboBox, QGroupBox, QSplitter,
                             QTableWidget, QTableWidgetItem, QHeaderView,
                             QAbstractItemView, QInputDialog, QMessageBox,
                             QFileDialog, QFrame, QMenu)

from app_core.i18n import I18n
from app_core.theme_engine import ThemeEngine
from app_core.utils import wrap_table_with_glow
from services.database import DatabaseManager
from services.email_service import EmailService
from services.telegram_bot import TelegramBot
from services.webhook_service import fire_event
from widgets.toast import ToastNotification


class RemindersDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.db = DatabaseManager()
        self.setWindowTitle(I18n._("reminder.all"))
        self.setMinimumSize(550, 450)
        self.resize(600, 500)
        self._build_ui()
        self._load_reminders()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(12)

        heading = QLabel(I18n._("reminder.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        self._table = QTableWidget()
        self._table.setColumnCount(5)
        self._table.setHorizontalHeaderLabels([
            I18n._("reminder.title_field"), I18n._("reminder.description"),
            I18n._("reminder.due_date"), I18n._("reminder.interval"),
            I18n._("reminder.is_done")])
        self._table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._table.setAlternatingRowColors(True)
        self._table.verticalHeader().hide()
        self._table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        layout.addWidget(wrap_table_with_glow(self._table, self))

        btn_layout = QHBoxLayout()
        self._add_btn = QPushButton(I18n._("reminder.add"))
        self._add_btn.setProperty("success", True)
        self._add_btn.clicked.connect(self._add_reminder)
        btn_layout.addWidget(self._add_btn)
        self._edit_btn = QPushButton(I18n._("common.edit"))
        self._edit_btn.clicked.connect(self._edit_reminder)
        btn_layout.addWidget(self._edit_btn)
        self._delete_btn = QPushButton(I18n._("reminder.delete"))
        self._delete_btn.setProperty("danger", True)
        self._delete_btn.clicked.connect(self._delete_reminder)
        btn_layout.addWidget(self._delete_btn)
        self._toggle_btn = QPushButton(I18n._("common.done"))
        self._toggle_btn.clicked.connect(self._toggle_done)
        btn_layout.addWidget(self._toggle_btn)
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)

    def _load_reminders(self) -> None:
        reminders = self.db.get_reminders(include_done=True)
        self._table.setRowCount(len(reminders))
        for i, r in enumerate(reminders):
            self._table.setItem(i, 0, QTableWidgetItem(r.get("title", "")))
            desc = r.get("description", "")
            self._table.setItem(i, 1, QTableWidgetItem(
                desc[:60] + ("..." if len(desc) > 60 else "")))
            self._table.setItem(i, 2, QTableWidgetItem(r.get("due_date", "")))
            self._table.setItem(i, 3, QTableWidgetItem(f"{r.get('check_interval', 60)}c"))
            done_item = QTableWidgetItem(
                "✓" if r.get("is_done") else "☐")
            done_item.setTextAlignment(Qt.AlignCenter)
            self._table.setItem(i, 4, done_item)
            self._table.item(i, 0).setData(Qt.UserRole, r["id"])
        self._table.resizeColumnsToContents()

    def _add_reminder(self) -> None:
        title, ok = QInputDialog.getText(self, I18n._("reminder.add"),
                                         I18n._("reminder.title_field"))
        if not ok or not title:
            return
        desc, ok2 = QInputDialog.getMultiLineText(self, I18n._("reminder.add"),
                                                   I18n._("reminder.description"))
        if not ok2:
            desc = ""
        due, ok3 = QInputDialog.getText(self, I18n._("reminder.add"),
                                         I18n._("reminder.due_date"),
                                         text=datetime.now().strftime("%d.%m.%Y"))
        if not ok3 or not due:
            return
        interval, ok4 = QInputDialog.getInt(self, I18n._("reminder.add"),
                                             I18n._("reminder.interval"), 60, 10, 86400)
        if not ok4:
            return
        self.db.save_reminder(title.strip(), desc.strip(), due.strip(), interval)
        self._load_reminders()
        ToastNotification.notify(I18n._("common.success"), "success", 3000)

    def _edit_reminder(self) -> None:
        row = self._table.currentRow()
        if row < 0:
            return
        rid = self._table.item(row, 0).data(Qt.UserRole)
        title = self._table.item(row, 0).text()
        desc = self._table.item(row, 1).text()
        due = self._table.item(row, 2).text()

        new_title, ok = QInputDialog.getText(self, I18n._("reminder.edit"),
                                              I18n._("reminder.title_field"),
                                              text=title)
        if not ok:
            return
        new_due, ok2 = QInputDialog.getText(self, I18n._("reminder.edit"),
                                             I18n._("reminder.due_date"),
                                             text=due)
        if not ok2:
            return
        self.db.save_reminder(new_title.strip(), desc, new_due.strip(),
                              60, reminder_id=rid)
        self._load_reminders()
        ToastNotification.notify(I18n._("common.success"), "success", 3000)

    def _delete_reminder(self) -> None:
        row = self._table.currentRow()
        if row < 0:
            return
        rid = self._table.item(row, 0).data(Qt.UserRole)
        reply = QMessageBox.question(self, I18n._("common.confirm"),
                                     I18n._("reminder.delete") + "?",
                                     QMessageBox.Yes | QMessageBox.No)
        if reply == QMessageBox.Yes:
            self.db.delete_reminder(rid)
            self._load_reminders()
            ToastNotification.notify(I18n._("toast.delete_success"), "success", 3000)

    def _toggle_done(self) -> None:
        row = self._table.currentRow()
        if row < 0:
            return
        rid = self._table.item(row, 0).data(Qt.UserRole)
        reminders = self.db.get_reminders(include_done=True)
        for r in reminders:
            if r["id"] == rid:
                new_done = 0 if r.get("is_done") else 1
                self.db.save_reminder(r["title"], r.get("description", ""),
                                      r["due_date"], r.get("check_interval", 60),
                                      is_done=new_done, reminder_id=rid)
                break
        self._load_reminders()


class ReminderEngine(QObject):
    def __init__(self, parent: Optional[QObject] = None) -> None:
        super().__init__(parent)
        self.db = DatabaseManager()
        self._notified: set = set()
        self._timer = QTimer(self)
        self._timer.timeout.connect(self._check)
        self._interval = int(self.db.get_setting("reminder_check_interval", "60"))
        self._timer.start(self._interval * 1000)

    def _check(self) -> None:
        try:
            reminders = self.db.get_due_reminders()
            now = datetime.now()
            for r in reminders:
                rid = r.get("id", 0)
                if rid in self._notified:
                    continue
                due_str = r.get("due_date", "")
                try:
                    p = due_str.split(".")
                    if len(p) == 3:
                        due = datetime(int(p[2]), int(p[1]), int(p[0]))
                        if due <= now and not r.get("is_done"):
                            self._notified.add(rid)
                            title = r.get("title", I18n._("reminder.title"))
                            ToastNotification.notify(
                                f"🔔 {title}", "warning", 8000)
                            try:
                                TelegramBot().send_notification(
                                    title,
                                    f"{I18n._('reminder.due_date')}: {r.get('due_date', '')}\n{r.get('description', '')}",
                                    "warning")
                            except Exception:
                                pass
                            try:
                                EmailService().send_notification(
                                    title,
                                    f"{I18n._('reminder.due_date')}: {r.get('due_date', '')}\n{r.get('description', '')}",
                                    "warning")
                            except Exception:
                                pass
                            try:
                                fire_event("reminder.due", {
                                    "id": rid,
                                    "title": title,
                                    "due_date": r.get("due_date", ""),
                                    "description": r.get("description", ""),
                                })
                            except Exception:
                                pass
                except Exception:
                    pass
            if reminders:
                active = {r["id"] for r in reminders}
                self._notified &= active
        except Exception:
            pass

    def update_interval(self, interval: int) -> None:
        self._interval = max(10, interval)
        self._timer.setInterval(self._interval * 1000)

    def stop(self) -> None:
        self._timer.stop()


class ExpiringRemindersTab(QWidget):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.db = DatabaseManager()
        self._quick_filter = self.db.get_setting("reminder_quick_filter", "all")
        self._build_ui()
        self._load_data()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(12)

        heading = QLabel(I18n._("reminder.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)

        hint = QLabel(I18n._("reminder.expiring_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet("font-size: 12px;")
        layout.addWidget(hint)

        quick_layout = QHBoxLayout()
        self._quick_buttons: Dict[str, QPushButton] = {}
        self._quick_button_labels: Dict[str, str] = {}
        for key, label_key in [
            ("all", "reminder.quick_all"),
            ("overdue", "reminder.quick_overdue"),
            ("3days", "reminder.quick_3days"),
            ("30days", "reminder.quick_30days"),
        ]:
            btn = QPushButton(I18n._(label_key))
            btn.setProperty("flat", True)
            btn.clicked.connect(lambda _=False, k=key: self._set_quick_filter(k))
            quick_layout.addWidget(btn)
            self._quick_buttons[key] = btn
            self._quick_button_labels[key] = label_key
        quick_layout.addStretch()
        layout.addLayout(quick_layout)

        filter_layout = QHBoxLayout()
        self._overdue_only_cb = QCheckBox(I18n._("reminder.overdue_only"))
        self._overdue_only_cb.toggled.connect(self._apply_filters)
        filter_layout.addWidget(self._overdue_only_cb)
        filter_layout.addSpacing(16)
        filter_layout.addWidget(QLabel(I18n._("reminder.sort_order") + ":"))
        self._sort_combo = QComboBox()
        self._sort_combo.addItem(I18n._("reminder.sort_priority"), "priority")
        self._sort_combo.addItem(I18n._("reminder.sort_due_asc"), "due_asc")
        self._sort_combo.addItem(I18n._("reminder.sort_due_desc"), "due_desc")
        saved_sort = self.db.get_setting("reminder_sort_order", "priority")
        idx = self._sort_combo.findData(saved_sort)
        if idx >= 0:
            self._sort_combo.setCurrentIndex(idx)
        self._sort_combo.currentIndexChanged.connect(self._on_sort_changed)
        filter_layout.addWidget(self._sort_combo)
        filter_layout.addStretch()
        layout.addLayout(filter_layout)

        splitter = QSplitter(Qt.Vertical)
        self._overdue_group = QGroupBox(I18n._("reminder.overdue_section"))
        overdue_layout = QVBoxLayout(self._overdue_group)
        self._overdue_table = self._create_table()
        self._overdue_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self._overdue_table.customContextMenuRequested.connect(lambda pos: self._show_table_menu(self._overdue_table, pos))
        self._overdue_table.itemDoubleClicked.connect(lambda item: self._open_item_from_table(self._overdue_table, item.row()))
        overdue_layout.addWidget(self._overdue_table)
        splitter.addWidget(self._overdue_group)

        self._upcoming_group = QGroupBox(I18n._("reminder.upcoming_section"))
        upcoming_layout = QVBoxLayout(self._upcoming_group)
        self._upcoming_table = self._create_table()
        self._upcoming_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self._upcoming_table.customContextMenuRequested.connect(lambda pos: self._show_table_menu(self._upcoming_table, pos))
        self._upcoming_table.itemDoubleClicked.connect(lambda item: self._open_item_from_table(self._upcoming_table, item.row()))
        upcoming_layout.addWidget(self._upcoming_table)
        splitter.addWidget(self._upcoming_group)
        splitter.setSizes([220, 300])
        layout.addWidget(splitter, 1)

        btn_layout = QHBoxLayout()
        self._refresh_btn = QPushButton(I18n._("common.refresh"))
        self._refresh_btn.clicked.connect(self._load_data)
        btn_layout.addWidget(self._refresh_btn)
        self._export_csv_btn = QPushButton(I18n._("common.export") + " CSV")
        self._export_csv_btn.clicked.connect(lambda: self._export_reminders("csv"))
        btn_layout.addWidget(self._export_csv_btn)
        self._export_excel_btn = QPushButton(I18n._("common.export") + " Excel")
        self._export_excel_btn.clicked.connect(lambda: self._export_reminders("xlsx"))
        btn_layout.addWidget(self._export_excel_btn)
        btn_layout.addStretch()
        self._open_btn = QPushButton(I18n._("reminder.all"))
        self._open_btn.clicked.connect(self._open_all)
        btn_layout.addWidget(self._open_btn)
        layout.addLayout(btn_layout)

        self._auto_refresh_timer = QTimer(self)
        self._auto_refresh_timer.timeout.connect(self._load_data)
        self._auto_refresh_timer.start(120000)

        self._update_quick_filter_buttons()

    def _create_table(self) -> QTableWidget:
        table = QTableWidget()
        table.setColumnCount(5)
        table.setHorizontalHeaderLabels([
            I18n._("reminder.type"), I18n._("reminder.title_field"),
            I18n._("reminder.entity"), I18n._("reminder.due_date"),
            I18n._("reminder.days_left")])
        header = table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setSectionResizeMode(0, QHeaderView.Stretch)
        table.setColumnWidth(1, 140)
        table.setColumnWidth(2, 180)
        table.setColumnWidth(3, 120)
        table.setColumnWidth(4, 100)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setAlternatingRowColors(True)
        table.verticalHeader().hide()
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSortingEnabled(False)
        return table

    def _set_quick_filter(self, key: str) -> None:
        self._quick_filter = key
        self.db.upsert_setting("reminder_quick_filter", key)
        self._update_quick_filter_buttons()
        self._apply_filters()

    def _update_quick_filter_buttons(self) -> None:
        accent = self.db.get_setting("accent_color", "#2196F3")
        items = list(getattr(self, "_items", []))
        counters = {
            "all": len(items),
            "overdue": len([x for x in items if x.get("days", 999) < 0]),
            "3days": len([x for x in items if 0 <= x.get("days", 999) <= 3]),
            "30days": len([x for x in items if 0 <= x.get("days", 999) <= 30]),
        }
        for key, btn in getattr(self, "_quick_buttons", {}).items():
            active = key == getattr(self, "_quick_filter", "all")
            label_key = self._quick_button_labels.get(key, "")
            btn.setText(f"{I18n._(label_key)} ({counters.get(key, 0)})")
            btn.setStyleSheet(
                f"padding: 6px 12px; border-radius: 6px;"
                f"background: {accent + ('30' if active else '10')};"
                f"border: 1px solid {accent if active else accent + '30'};"
                f"font-weight: {'600' if active else '500'};")

    def _on_sort_changed(self) -> None:
        self.db.upsert_setting("reminder_sort_order", self._sort_combo.currentData())
        self._apply_filters()

    def _load_data(self) -> None:
        self._items: List[Dict[str, Any]] = []
        now = datetime.now()
        try:
            employees = self.db.get_json_records("employees")
            for emp in employees:
                dj = emp.get("data_json", {})
                for key in ("Дата медосмотра",):
                    val = dj.get(key, "")
                    try:
                        p = val.split(".")
                        if len(p) == 3:
                            dt = datetime(int(p[2]), int(p[1]), int(p[0]))
                            left = (dt - now).days
                            self._items.append({
                                "type": I18n._("tab.employees"),
                                "title": key,
                                "entity": dj.get("ФИО", f"#{emp['id']}"),
                                "due": val,
                                "days": left,
                                "_table": "employees",
                                "_record_id": emp.get("id"),
                                "_ts": dt.timestamp(),
                            })
                    except Exception:
                        pass

            violations = self.db.get_json_records("violations")
            for viol in violations:
                dj = viol.get("data_json", {})
                for key in ("Срок устранения", "Дата", "Годен до"):
                    val = dj.get(key, "")
                    try:
                        p = val.split(".")
                        if len(p) == 3:
                            dt = datetime(int(p[2]), int(p[1]), int(p[0]))
                            left = (dt - now).days
                            self._items.append({
                                "type": I18n._("tab.violations"),
                                "title": key,
                                "entity": f"#{viol['id']} {dj.get('Описание', '')[:30]}",
                                "due": val,
                                "days": left,
                                "_table": "violations",
                                "_record_id": viol.get("id"),
                                "_ts": dt.timestamp(),
                            })
                    except Exception:
                        pass

            reminders = self.db.get_reminders(include_done=False)
            for r in reminders:
                due_str = r.get("due_date", "")
                try:
                    p = due_str.split(".")
                    if len(p) == 3:
                        dt = datetime(int(p[2]), int(p[1]), int(p[0]))
                        left = (dt - now).days
                        self._items.append({
                            "type": I18n._("reminder.title"),
                            "title": r.get("title", ""),
                            "entity": "",
                            "due": due_str,
                            "days": left,
                            "_table": "reminders",
                            "_record_id": r.get("id"),
                            "_id": r["id"],
                            "_done": r.get("is_done", False),
                            "_ts": dt.timestamp(),
                        })
                except Exception:
                    pass
        except Exception:
            import traceback; traceback.print_exc()

        self._apply_filters()

    def _apply_filters(self) -> None:
        items = list(getattr(self, "_items", []))
        quick = getattr(self, "_quick_filter", "all")
        if quick == "overdue":
            items = [item for item in items if item.get("days", 999) < 0]
        elif quick == "3days":
            items = [item for item in items if 0 <= item.get("days", 999) <= 3]
        elif quick == "30days":
            items = [item for item in items if 0 <= item.get("days", 999) <= 30]

        if getattr(self, "_overdue_only_cb", None) and self._overdue_only_cb.isChecked():
            items = [item for item in items if item.get("days", 999) < 0]

        sort_mode = self._sort_combo.currentData() if hasattr(self, "_sort_combo") else "priority"

        def _sort_key(item: Dict[str, Any]) -> Tuple[int, int, float]:
            d = int(item.get("days", 999))
            ts = float(item.get("_ts", 0.0))
            if sort_mode == "due_desc":
                return (0, 0, -ts)
            if sort_mode == "due_asc":
                return (0, 0, ts)
            group = 0 if d < 0 else (1 if d <= 3 else 2)
            return (group, abs(d) if d < 0 else d, ts)

        items.sort(key=_sort_key)
        self._visible_items = items
        overdue_items = [item for item in items if item.get("days", 999) < 0]
        upcoming_items = [item for item in items if item.get("days", 999) >= 0]
        self._fill_table(self._overdue_table, overdue_items)
        self._fill_table(self._upcoming_table, upcoming_items)
        self._overdue_group.setTitle(f"{I18n._('reminder.overdue_section')} ({len(overdue_items)})")
        self._upcoming_group.setTitle(f"{I18n._('reminder.upcoming_section')} ({len(upcoming_items)})")
        self._overdue_group.setVisible(bool(overdue_items) or quick in ("all", "overdue") or self._overdue_only_cb.isChecked())
        self._upcoming_group.setVisible(not self._overdue_only_cb.isChecked())
        self._update_quick_filter_buttons()

    def _fill_table(self, table: QTableWidget, items: List[Dict[str, Any]]) -> None:
        table.setSortingEnabled(False)
        table.setRowCount(len(items))
        icon_map = {
            I18n._("tab.employees"): "👤 ",
            I18n._("tab.violations"): "⚠ ",
            I18n._("reminder.title"): "🔔 ",
        }
        is_dark = ThemeEngine._current_theme == "dark"
        for i, item in enumerate(items):
            prefix = icon_map.get(item["type"], "")
            days = item["days"]
            if days < 0:
                days_str = I18n._("reminder.overdue").format(days=abs(days))
                row_bg = QColor("#582525" if is_dark else "#f8d7da")
            elif days <= 3:
                days_str = I18n._("reminder.today") if days == 0 else I18n._("reminder.days_format").format(days=days)
                row_bg = QColor("#614d17" if is_dark else "#fff3cd")
            else:
                days_str = I18n._("reminder.days_format").format(days=days)
                row_bg = QColor("#254b32" if is_dark else "#d4edda")
            row_fg = QColor("#FFFFFF") if row_bg.lightness() < 140 else QColor("#1E1E2E")
            values = [prefix + item["type"], item["title"], item["entity"], item["due"], days_str]
            for col, value in enumerate(values):
                cell = QTableWidgetItem(value)
                cell.setBackground(row_bg)
                cell.setForeground(row_fg)
                if col == 4:
                    cell.setTextAlignment(Qt.AlignCenter)
                cell.setData(Qt.UserRole, item)
                table.setItem(i, col, cell)
        table.resizeColumnsToContents()
        table.setSortingEnabled(True)

    def _open_item_from_table(self, table: QTableWidget, row: int) -> None:
        item = table.item(row, 0)
        if not item:
            return
        data = item.data(Qt.UserRole) or {}
        mw = self.window()
        if mw and hasattr(mw, "open_reminder_target"):
            mw.open_reminder_target(data)

    def _show_table_menu(self, table: QTableWidget, pos: QPoint) -> None:
        row = table.rowAt(pos.y())
        if row >= 0:
            table.selectRow(row)
        menu = QMenu(self)
        open_a = menu.addAction(I18n._("common.edit"))
        refresh_a = menu.addAction(I18n._("common.refresh"))
        export_csv_a = menu.addAction(I18n._("common.export") + " CSV")
        export_xlsx_a = menu.addAction(I18n._("common.export") + " Excel")
        action = menu.exec_(table.mapToGlobal(pos))
        if action == open_a and row >= 0:
            self._open_item_from_table(table, row)
        elif action == refresh_a:
            self._load_data()
        elif action == export_csv_a:
            self._export_reminders("csv")
        elif action == export_xlsx_a:
            self._export_reminders("xlsx")

    def _export_reminders(self, fmt: str = "csv") -> None:
        items = list(getattr(self, "_visible_items", []))
        if not items:
            return
        if fmt == "csv":
            path, _ = QFileDialog.getSaveFileName(self, I18n._("export.title"),
                                                  "reminders.csv", "CSV (*.csv)")
            if not path:
                return
            try:
                with open(path, "w", newline="", encoding="utf-8-sig") as f:
                    w = csv.writer(f)
                    w.writerow([I18n._("reminder.type"), I18n._("reminder.title_field"),
                                I18n._("reminder.entity"), I18n._("reminder.due_date"),
                                I18n._("reminder.days_left")])
                    for item in items:
                        days = item.get("days", 0)
                        days_str = I18n._("reminder.overdue").format(days=abs(days)) if days < 0 else (
                            I18n._("reminder.today") if days == 0 else I18n._("reminder.days_format").format(days=days))
                        w.writerow([item.get("type", ""), item.get("title", ""), item.get("entity", ""), item.get("due", ""), days_str])
                ToastNotification.notify(I18n._("export.success").format(path=path), "success", 3000)
            except Exception as e:
                ToastNotification.notify(I18n._("export.error").format(error=str(e)), "error", 5000)
            return
        path, _ = QFileDialog.getSaveFileName(self, I18n._("export.title"),
                                              "reminders.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reminders"
            headers = [I18n._("reminder.type"), I18n._("reminder.title_field"),
                       I18n._("reminder.entity"), I18n._("reminder.due_date"),
                       I18n._("reminder.days_left")]
            for c, h in enumerate(headers, 1):
                ws.cell(row=1, column=c, value=h)
            for r, item in enumerate(items, 2):
                days = item.get("days", 0)
                days_str = I18n._("reminder.overdue").format(days=abs(days)) if days < 0 else (
                    I18n._("reminder.today") if days == 0 else I18n._("reminder.days_format").format(days=days))
                vals = [item.get("type", ""), item.get("title", ""), item.get("entity", ""), item.get("due", ""), days_str]
                for c, v in enumerate(vals, 1):
                    ws.cell(row=r, column=c, value=v)
            wb.save(path)
            ToastNotification.notify(I18n._("export.success").format(path=path), "success", 3000)
        except Exception as e:
            ToastNotification.notify(I18n._("export.error").format(error=str(e)), "error", 5000)

    def has_expiring(self) -> bool:
        return any(item.get("days", 999) <= 3 for item in getattr(self, "_items", []))

    def get_expiring_count(self) -> int:
        return len([item for item in getattr(self, "_items", []) if item.get("days", 999) <= 30])

    def _open_all(self) -> None:
        RemindersDialog(self).exec_()

    def refresh(self) -> None:
        self._load_data()
        self._update_quick_filter_buttons()


class ReminderFloatingDialog(QDialog):
    def __init__(self, parent: Optional[QWidget] = None) -> None:
        super().__init__(parent)
        self.setWindowFlags(
            Qt.Window | Qt.WindowStaysOnTopHint | Qt.WindowCloseButtonHint |
            Qt.WindowMinimizeButtonHint | Qt.WindowMaximizeButtonHint)
        self.setWindowTitle(I18n._("reminder.title"))
        self.setMinimumSize(500, 350)
        self.resize(600, 400)
        self.db = DatabaseManager()
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(12)
        heading = QLabel(I18n._("reminder.title"))
        heading.setProperty("heading", True)
        layout.addWidget(heading)
        self._table = QTableWidget()
        self._table.setColumnCount(5)
        self._table.setHorizontalHeaderLabels([
            I18n._("reminder.type"), I18n._("reminder.title_field"),
            I18n._("reminder.entity"), I18n._("reminder.due_date"),
            I18n._("reminder.days_left")])
        hdr = self._table.horizontalHeader()
        hdr.setSectionResizeMode(QHeaderView.Interactive)
        hdr.setStretchLastSection(True)
        self._table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._table.setAlternatingRowColors(True)
        self._table.verticalHeader().hide()
        self._table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        layout.addWidget(self._table, 1)
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()
        close_btn = QPushButton(I18n._("common.close"))
        close_btn.clicked.connect(self.close)
        btn_layout.addWidget(close_btn)
        layout.addLayout(btn_layout)
        self._load_data()

    def _load_data(self) -> None:
        now = datetime.now()
        items: List[Tuple[str, str, str, str, int, float]] = []
        try:
            for emp in self.db.get_json_records("employees"):
                dj = emp.get("data_json", {})
                for key in ("Дата медосмотра",):
                    val = dj.get(key, "")
                    try:
                        p = val.split(".")
                        if len(p) == 3:
                            dt = datetime(int(p[2]), int(p[1]), int(p[0]))
                            left = (dt - now).days
                            if left <= 30:
                                items.append(("👤 " + I18n._("tab.employees"), key,
                                              dj.get("ФИО", f"#{emp['id']}"), val, left, dt.timestamp()))
                    except Exception:
                        pass
            for viol in self.db.get_json_records("violations"):
                dj = viol.get("data_json", {})
                for key in ("Срок устранения", "Дата", "Годен до"):
                    val = dj.get(key, "")
                    try:
                        p = val.split(".")
                        if len(p) == 3:
                            dt = datetime(int(p[2]), int(p[1]), int(p[0]))
                            left = (dt - now).days
                            if left <= 30:
                                items.append(("⚠ " + I18n._("tab.violations"), key,
                                              f"#{viol['id']}", val, left, dt.timestamp()))
                    except Exception:
                        pass
        except Exception:
            pass
        items.sort(key=lambda x: (0 if x[4] < 0 else 1, x[5]))
        self._table.setRowCount(len(items))
        for i, (typ, title, entity, due, days, _ts) in enumerate(items):
            if days < 0:
                ds = I18n._("reminder.overdue").format(days=abs(days))
            elif days == 0:
                ds = I18n._("reminder.today")
            else:
                ds = I18n._("reminder.days_format").format(days=days)
            self._table.setItem(i, 0, QTableWidgetItem(typ))
            self._table.setItem(i, 1, QTableWidgetItem(title))
            self._table.setItem(i, 2, QTableWidgetItem(entity))
            self._table.setItem(i, 3, QTableWidgetItem(due))
            di = QTableWidgetItem(ds)
            di.setTextAlignment(Qt.AlignCenter)
            if days < 0:
                di.setForeground(QColor("#E53935"))
            elif days <= 3:
                di.setForeground(QColor("#FF9800"))
            else:
                di.setForeground(QColor("#4CAF50"))
            self._table.setItem(i, 4, di)
        self._table.resizeColumnsToContents()

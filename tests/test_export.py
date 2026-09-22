"""Часть 10: тесты экспорта (xlsx/csv/json/docx/zip/bundle, диапазоны)."""

import io
import os
import sys
import tempfile
import zipfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

p = os.path.join(tempfile.gettempdir(), "suot_exp.db")
for s in ("", "-wal", "-shm"):
    try:
        os.remove(p + s)
    except OSError:
        pass
os.environ["SUOT_E2E_DB"] = p

from fastapi.testclient import TestClient  # noqa: E402
from server.app import app  # noqa: E402

PASS, FAIL = [], []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("  OK " if cond else "FAIL ") + name + (f" | {extra}" if extra else ""))


client = TestClient(app)
with client:
    r = client.post(
        "/api/auth/register", json={"username": "exp_user", "password": "parol123"}
    )
    HB = {"Authorization": "Bearer " + r.json()["token"]}

    # данные
    for i, (fio, dep) in enumerate(
        [("Алексеев А.", "Цех №1"), ("Борисов Б.", "Цех №2"), ("Смирнова В.", "Склад")]
    ):
        client.post(
            "/api/data/employees",
            headers=HB,
            json={"data": {"ФИО": fio, "Подразделение": dep}},
        )

    body = {"table": "employees"}

    # xlsx
    rx = client.post("/api/export/xlsx/employees", headers=HB, json=body)
    check(
        "xlsx: 200 и размер",
        rx.status_code == 200 and len(rx.content) > 2000,
        str(len(rx.content)),
    )
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(rx.content))
    ws = wb.active
    hdrs = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    check("xlsx: заголовки", "ФИО" in hdrs, str(hdrs[:4]))
    check("xlsx: 3 строки данных", ws.max_row == 4, str(ws.max_row))

    # выбранные ids
    ids = [
        i["id"]
        for i in client.get("/api/data/employees", headers=HB).json()["items"][:1]
    ]
    rx2 = client.post(
        "/api/export/xlsx/employees", headers=HB, json={**body, "ids": ids}
    )
    wb2 = openpyxl.load_workbook(io.BytesIO(rx2.content))
    check(
        "xlsx: выбранные строки (1)", wb2.active.max_row == 2, str(wb2.active.max_row)
    )

    # выбор колонок
    rx3 = client.post(
        "/api/export/xlsx/employees",
        headers=HB,
        json={"table": "employees", "columns": ["ФИО"]},
    )
    wb3 = openpyxl.load_workbook(io.BytesIO(rx3.content))
    check(
        "xlsx: выбор колонок (только ФИО)",
        wb3.active.max_column == 1,
        str(wb3.active.max_column),
    )

    # фильтр как в UI (q)
    rx4 = client.post("/api/export/xlsx/employees?query=1", headers=HB, json=body)
    check("экспорт с query-параметрами не падает", rx4.status_code == 200)

    # csv
    rc = client.post("/api/export/csv/employees", headers=HB, json=body)
    check(
        "csv: 200, BOM, разделитель ;",
        rc.status_code == 200
        and rc.content[:3] == b"\xef\xbb\xbf"
        and b";" in rc.content[:200],
    )

    # json
    rj = client.post("/api/export/json/employees", headers=HB, json=body)
    jj = rj.json()
    check("json: 3 записи", rj.status_code == 200 and len(jj["items"]) == 3)

    # docx
    rd = client.post("/api/export/docx/employees", headers=HB, json=body)
    check(
        "docx: 200",
        rd.status_code == 200 and len(rd.content) > 5000,
        str(len(rd.content)),
    )

    # своя колонка попадает в экспорт
    ruc = client.post(
        "/api/ucols/employees",
        headers=HB,
        json={"name": "Е2Е Колонка", "type": "Текст"},
    )
    own = client.get("/api/data/employees?page_size=1", headers=HB).json()["items"][0][
        "id"
    ]
    ed = client.get(f"/api/data/employees/{own}", headers=HB).json()["data"]
    ed["Е2Е Колонка"] = "Значение"
    client.put(f"/api/data/employees/{own}", headers=HB, json={"data": ed})
    ruc_e = client.post("/api/export/csv/employees", headers=HB, json=body)
    check("своя колонка в экспорте", "Е2Е Колонка".encode("utf-8") in ruc_e.content)

    # чужая таблица недоступна
    r2 = client.post(
        "/api/auth/register", json={"username": "exp2", "password": "parol123"}
    )
    HB2 = {"Authorization": "Bearer " + r2.json()["token"]}
    rt = client.post(
        "/api/custom/tables",
        headers=HB,
        json={"label": "Приватная", "columns": [{"name": "Название"}]},
    )
    rden = client.post(
        f"/api/export/csv/{rt.json()['key']}",
        headers=HB2,
        json={"table": rt.json()["key"]},
    )
    check("чужая u_ таблица → 403", rden.status_code == 403)

    # bundle
    rb = client.post("/api/export/bundle", headers=HB)
    check(
        "bundle: ZIP с таблицами",
        rb.status_code == 200
        and rb.content[:2] == b"PK"
        and b"employees.xlsx" in rb.content,
    )

for s in ("", "-wal", "-shm"):
    try:
        os.remove(p + s)
    except OSError:
        pass

print(f"\nИТОГО: {len(PASS)} OK, {len(FAIL)} FAIL")
if FAIL:
    print("Провалены:", FAIL)
    sys.exit(1)

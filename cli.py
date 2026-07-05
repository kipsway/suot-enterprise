#!/usr/bin/env python3
"""
SUOT Enterprise CLI — headless operation mode.

Usage:
    python cli.py violations list [--limit N]
    python cli.py violations export [--format csv|xlsx] [--output FILE]
    python cli.py employees list [--limit N]
    python cli.py employees export [--format csv|xlsx] [--output FILE]
    python cli.py companies list
    python cli.py reminders check
    python cli.py stats
    python cli.py backup create
    python cli.py backup list
    python cli.py server [--port PORT]
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime
from typing import Any, Dict, List, Optional

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app_core.config import RUNTIME_PATHS
from app_core.i18n import I18n
from services.database import DatabaseManager


def setup() -> DatabaseManager:
    os.makedirs(RUNTIME_PATHS.backup_dir, exist_ok=True)
    os.makedirs(RUNTIME_PATHS.media_dir, exist_ok=True)
    db = DatabaseManager()
    lang = db.get_setting("app_language", "ru")
    I18n.set_language(lang)
    return db


def cmd_violations_list(args: argparse.Namespace) -> None:
    db = setup()
    records = db.get_json_records("violations", limit=args.limit)
    print(f"Violations: {len(records)}")
    print(f"{'ID':>4} {'Status':<12} {'Company':<20} {'Description':<40} {'Deadline':<12}")
    print("-" * 92)
    for r in records:
        dj = r.get("data_json", r)
        rid = r["id"]
        status = dj.get("Статус", "?")
        company = str(dj.get("Фирма", ""))[:20]
        desc = str(dj.get("Описание", ""))[:40]
        deadline = dj.get("Срок устранения", "")
        print(f"{rid:>4} {status:<12} {company:<20} {desc:<40} {deadline:<12}")


def cmd_violations_export(args: argparse.Namespace) -> None:
    db = setup()
    records = db.get_json_records("violations", limit=5000)
    _export_table(records, "violations", args)


def cmd_employees_list(args: argparse.Namespace) -> None:
    db = setup()
    records = db.get_json_records("employees", limit=args.limit)
    print(f"Employees: {len(records)}")
    print(f"{'ID':>4} {'Name':<30} {'Company':<20} {'Position':<25}")
    print("-" * 83)
    for r in records:
        dj = r.get("data_json", r)
        name = str(dj.get("ФИО", f"#{r['id']}"))[:30]
        company = str(dj.get("Фирма", ""))[:20]
        position = str(dj.get("Должность", ""))[:25]
        print(f"{r['id']:>4} {name:<30} {company:<20} {position:<25}")


def cmd_employees_export(args: argparse.Namespace) -> None:
    db = setup()
    records = db.get_json_records("employees", limit=5000)
    _export_table(records, "employees", args)


def cmd_companies_list(args: argparse.Namespace) -> None:
    db = setup()
    companies = db.get_companies()
    print(f"Companies: {len(companies)}")
    print(f"{'ID':>4} {'Name':<30} {'Address':<30}")
    print("-" * 68)
    for c in companies:
        name = str(c.get("name", ""))[:30]
        addr = str(c.get("address", ""))[:30]
        print(f"{c['id']:>4} {name:<30} {addr:<30}")


def cmd_reminders_check(args: argparse.Namespace) -> None:
    db = setup()
    reminders = db.get_due_reminders()
    now = datetime.now()
    overdue = []
    upcoming = []
    for r in reminders:
        due_str = r.get("due_date", "")
        try:
            p = due_str.split(".")
            if len(p) == 3:
                due = datetime(int(p[2]), int(p[1]), int(p[0]))
                if due <= now:
                    overdue.append(r)
                else:
                    upcoming.append(r)
        except Exception:
            pass

    if overdue:
        print(f"\nOVERDUE REMINDERS ({len(overdue)}):")
        for r in overdue:
            print(f"  [{r.get('due_date', '')}] {r.get('title', '?')}")
    if upcoming:
        print(f"\nUPCOMING REMINDERS ({len(upcoming)}):")
        for r in upcoming[:10]:
            print(f"  [{r.get('due_date', '')}] {r.get('title', '?')}")
    if not overdue and not upcoming:
        print("No due reminders.")


def cmd_stats(args: argparse.Namespace) -> None:
    db = setup()
    stats = db.get_statistics()
    print("=" * 40)
    print("  SUOT Enterprise — Statistics")
    print("=" * 40)
    print(f"  Employees:  {stats['employees_total']}")
    print(f"  Violations: {stats['violations_total']}")
    print(f"  Companies:  {stats['companies_total']}")
    print(f"  Overdue:    {stats['overdue_total']}")
    print(f"  Total fines: {stats['fines_total']:,.0f}")

    records = db.get_json_records("violations", limit=5000)
    statuses: Dict[str, int] = {}
    for r in records:
        s = r.get("data_json", r).get("Статус", "?")
        statuses[s] = statuses.get(s, 0) + 1
    if statuses:
        print(f"\n  By status:")
        for s, c in sorted(statuses.items(), key=lambda x: -x[1]):
            print(f"    {s}: {c}")


def cmd_backup_create(args: argparse.Namespace) -> None:
    db = setup()
    path = db.create_backup()
    print(f"Backup created: {path}")


def cmd_backup_list(args: argparse.Namespace) -> None:
    db = setup()
    backups = db.get_backups()
    print(f"{'ID':>4} {'Date':<20} {'Size':<12} {'File':<50}")
    print("-" * 90)
    for b in backups:
        bid = b["id"]
        dt = b["created_at"][:19] if b["created_at"] else "?"
        size = f"{b['size_bytes'] / 1024:.0f} KB"
        fname = os.path.basename(b["file_path"])
        print(f"{bid:>4} {dt:<20} {size:<12} {fname:<50}")


def cmd_server(args: argparse.Namespace) -> None:
    from services.rest_api import RESTAPIServer, APIHandler
    db = setup()
    port = args.port
    api_key = db.get_setting("rest_api_key", "")
    APIHandler.api_key = api_key
    from http.server import HTTPServer
    server = HTTPServer(("127.0.0.1", port), APIHandler)
    print(f"REST API server started on http://127.0.0.1:{port}")
    if api_key:
        print(f"Auth: X-API-Key: {api_key[:8]}...")
    else:
        print("WARNING: No API key set — no authentication")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nShutting down...")
        server.shutdown()


def _export_table(records: List[Dict[str, Any]], label: str,
                  args: argparse.Namespace) -> None:
    fmt = args.format or "csv"
    out = args.output or f"{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{fmt}"
    flat = []
    for r in records:
        dj = r.get("data_json", r)
        flat.append({k: v for k, v in dj.items() if not k.startswith("_")})

    if fmt == "csv":
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            if flat:
                w = csv.DictWriter(f, fieldnames=list(flat[0].keys()))
                w.writeheader()
                w.writerows(flat)
        print(f"Exported {len(flat)} records to {out}")

    elif fmt == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = label
            if flat:
                headers = list(flat[0].keys())
                for c, h in enumerate(headers, 1):
                    ws.cell(row=1, column=c, value=h)
                for ri, row_data in enumerate(flat, 2):
                    for ci, h in enumerate(headers, 1):
                        ws.cell(row=ri, column=ci, value=str(row_data.get(h, "")))
            wb.save(out)
            print(f"Exported {len(flat)} records to {out}")
        except ImportError:
            print("openpyxl not installed. Use --format csv")
    else:
        print(f"Unsupported format: {fmt}. Use csv or xlsx.")


def main() -> None:
    parser = argparse.ArgumentParser(description="SUOT Enterprise CLI")
    sub = parser.add_subparsers(dest="command")

    # violations
    v = sub.add_parser("violations")
    v_sub = v.add_subparsers(dest="subcommand")
    v_list = v_sub.add_parser("list")
    v_list.add_argument("--limit", type=int, default=50)
    v_export = v_sub.add_parser("export")
    v_export.add_argument("--format", choices=["csv", "xlsx"], default="csv")
    v_export.add_argument("--output", "-o", type=str, default="")

    # employees
    e = sub.add_parser("employees")
    e_sub = e.add_subparsers(dest="subcommand")
    e_list = e_sub.add_parser("list")
    e_list.add_argument("--limit", type=int, default=50)
    e_export = e_sub.add_parser("export")
    e_export.add_argument("--format", choices=["csv", "xlsx"], default="csv")
    e_export.add_argument("--output", "-o", type=str, default="")

    # companies
    c = sub.add_parser("companies")
    c_sub = c.add_subparsers(dest="subcommand")
    c_list = c_sub.add_parser("list")

    # reminders
    r = sub.add_parser("reminders")
    r_sub = r.add_subparsers(dest="subcommand")
    r_sub.add_parser("check")

    # stats
    sub.add_parser("stats")

    # backup
    b = sub.add_parser("backup")
    b_sub = b.add_subparsers(dest="subcommand")
    b_sub.add_parser("create")
    b_sub.add_parser("list")

    # server
    s = sub.add_parser("server")
    s.add_argument("--port", type=int, default=8888)

    args = parser.parse_args()

    if args.command == "violations":
        if args.subcommand == "list":
            cmd_violations_list(args)
        elif args.subcommand == "export":
            cmd_violations_export(args)
        else:
            parser.print_help()
    elif args.command == "employees":
        if args.subcommand == "list":
            cmd_employees_list(args)
        elif args.subcommand == "export":
            cmd_employees_export(args)
        else:
            parser.print_help()
    elif args.command == "companies":
        cmd_companies_list(args)
    elif args.command == "reminders":
        cmd_reminders_check(args)
    elif args.command == "stats":
        cmd_stats(args)
    elif args.command == "backup":
        if args.subcommand == "create":
            cmd_backup_create(args)
        elif args.subcommand == "list":
            cmd_backup_list(args)
        else:
            parser.print_help()
    elif args.command == "server":
        cmd_server(args)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()

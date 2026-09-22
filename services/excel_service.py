"""Excel (XLSX) export with formatting."""

import os
from datetime import datetime
from typing import Any, Dict, List, Optional


def is_available() -> bool:
    try:
        import openpyxl

        return True
    except ImportError:
        return False


def _auto_width(ws, headers: List[str], data: List[Dict[str, Any]]) -> None:
    for ci, h in enumerate(headers, 1):
        max_len = len(str(h))
        for row in data:
            val = str(row.get(h, ""))
            max_len = max(max_len, len(val))
        ws.column_dimensions[chr(64 + ci) if ci <= 26 else "A"].width = min(
            max_len + 4, 60
        )


def export_to_excel(
    data: List[Dict[str, Any]],
    headers: List[str],
    output_path: str,
    sheet_name: str = "Sheet1",
    title: Optional[str] = None,
) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]
        header_fill = PatternFill(
            start_color="007AFF", end_color="007AFF", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True, size=12, name="Segoe UI")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thin", color="D0D0D0"),
        )
        row_num = 1
        if title:
            ws.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=len(headers)
            )
            cell = ws.cell(row=1, column=1, value=title)
            cell.font = Font(bold=True, size=16, name="Segoe UI", color="1C1C1E")
            cell.alignment = Alignment(horizontal="center")
            row_num = 2
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=ci, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border
        row_num += 1
        alt_fill = PatternFill(
            start_color="F2F2F7", end_color="F2F2F7", fill_type="solid"
        )
        for ri, record in enumerate(data):
            for ci, h in enumerate(headers, 1):
                val = record.get(h, "")
                cell = ws.cell(row=row_num, column=ci, value=str(val))
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if ri % 2 == 1:
                    cell.fill = alt_fill
            row_num += 1
        _auto_width(ws, headers, data)
        wb.save(output_path)
        return True
    except ImportError:
        return False


def export_table(
    table_name: str, output_path: str, title: Optional[str] = None
) -> bool:
    from services.database import DatabaseManager

    db = DatabaseManager()
    records = db.get_json_records(table_name, limit=10000)
    if not records:
        return False
    flat = []
    for r in records:
        dj = r.get("data_json", r)
        flat.append({k: v for k, v in dj.items() if not k.startswith("_")})
    headers = list(dict.fromkeys(k for row in flat for k in row))
    return export_to_excel(
        flat, headers, output_path, sheet_name=table_name, title=title or table_name
    )

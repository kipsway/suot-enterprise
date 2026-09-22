"""DOCX import/export service."""

import os
import tempfile
from typing import Any, Dict, List, Optional


def is_available() -> bool:
    try:
        import docx

        return True
    except ImportError:
        return False


def html_to_docx(html: str, output_path: str) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from bs4 import BeautifulSoup

        doc = Document()
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup.find_all(
            ["p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "table", "hr"]
        ):
            tag_name = tag.name
            if tag_name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                level = int(tag_name[1])
                heading = doc.add_heading(tag.get_text(strip=True), level=level)
            elif tag_name == "p":
                p = doc.add_paragraph(tag.get_text(strip=True))
            elif tag_name == "li":
                doc.add_paragraph(tag.get_text(strip=True), style="List Bullet")
            elif tag_name == "hr":
                doc.add_paragraph("_" * 60)
            elif tag_name == "table":
                rows = tag.find_all("tr")
                if rows:
                    table = doc.add_table(
                        rows=len(rows), cols=len(rows[0].find_all(["td", "th"]))
                    )
                    for ri, row in enumerate(rows):
                        cells = row.find_all(["td", "th"])
                        for ci, cell in enumerate(cells):
                            if ci < len(table.rows[ri].cells):
                                table.rows[ri].cells[ci].text = cell.get_text(
                                    strip=True
                                )
        doc.save(output_path)
        return True
    except ImportError:
        return False


def docx_to_html(docx_path: str) -> str:
    try:
        from docx import Document

        doc = Document(docx_path)
        parts = ["<!DOCTYPE html><html><body>"]
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading"):
                level = p.style.name.replace("Heading ", "")
                parts.append(f"<h{level}>{p.text}</h{level}>")
            else:
                parts.append(f"<p>{p.text}</p>")
        if doc.tables:
            for table in doc.tables:
                parts.append("<table>")
                for row in table.rows:
                    parts.append("<tr>")
                    for cell in row.cells:
                        parts.append(f"<td>{cell.text}</td>")
                    parts.append("</tr>")
                parts.append("</table>")
        parts.append("</body></html>")
        return "\n".join(parts)
    except ImportError:
        return ""


def export_record_to_docx(
    data: Dict[str, Any], columns: List[str], output_path: str, title: str = "Record"
) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for col in columns:
            val = str(data.get(col, ""))
            p = doc.add_paragraph()
            run = p.add_run(f"{col}: ")
            run.bold = True
            p.add_run(val)
        doc.save(output_path)
        return True
    except ImportError:
        return False


def export_records_to_docx(
    records: List[Dict[str, Any]],
    columns: List[str],
    output_path: str,
    title: str = "Report",
) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Total records: {len(records)}")
        doc.add_paragraph("")
        for i, record in enumerate(records, 1):
            doc.add_heading(f"Record #{i}", level=2)
            for col in columns:
                val = str(record.get(col, ""))
                p = doc.add_paragraph()
                run = p.add_run(f"{col}: ")
                run.bold = True
                p.add_run(val)
            doc.add_paragraph("_" * 60)
        doc.save(output_path)
        return True
    except ImportError:
        return False

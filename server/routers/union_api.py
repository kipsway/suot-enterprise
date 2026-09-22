"""Сводный реестр «Всё» (Часть 24): UNION всех JSON_TABLES + пользовательских
таблиц с колонкой «Раздел», сквозные q / фильтры по разделам / сортировка /
экспорт. Каждая запись помечается ключом и меткой раздела."""

import io
from datetime import datetime
import json
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from server.deps import get_db, get_current_user, is_admin
from services.database import DatabaseManager
from app_core.utils import JsonUtils

router = APIRouter(prefix="/api/union", tags=["union"])

TABLE_LABELS_RU = {
    "employees": "Сотрудники",
    "violations": "Нарушения",
    "custom_ledger": "Реестр",
    "incidents": "Происшествия",
    "ppe": "СИЗ",
    "training": "Обучение",
    "permits": "Допуски",
    "work_orders": "Наряды",
    "ppe_inspections": "Осмотры СИЗ",
    "companies": "Компании",
}

MAX_PER_SECTION = 5000


def _label(table: str) -> str:
    return TABLE_LABELS_RU.get(table, table)


def _sections(db: DatabaseManager, user: dict) -> List[Dict[str, Any]]:
    """Все доступные разделы: системные JSON_TABLES + custom (u_*)."""
    uid = None if is_admin(user) else int(user["id"])
    admin = is_admin(user)
    out: List[Dict[str, Any]] = []
    for table in DatabaseManager.JSON_TABLES:
        try:
            _, total = db.query_json_records(
                table, owner_id=uid, is_admin=admin, page_size=1
            )
        except Exception:
            total = 0
        cols = [c["name"] for c in db.get_columns_config(table) if c["name"] != "ID"]
        out.append(
            {
                "section": table,
                "label": _label(table),
                "icon": "",
                "count": int(total),
                "columns": cols,
            }
        )
    # пользовательские таблицы (не в корзине)
    for t in db.get_custom_tables(uid, admin):
        try:
            cols = [
                c["name"]
                for c in JsonUtils.loads(t["columns_json"] or "[]")
                if c.get("name") != "ID"
            ]
        except Exception:
            cols = []
        try:
            _, total = db.query_custom_records(
                t["key"], owner_id=uid, is_admin=admin, page_size=1
            )
        except Exception:
            total = 0
        out.append(
            {
                "section": t["key"],
                "label": t["label"],
                "icon": t.get("icon", "database"),
                "count": int(total),
                "columns": cols,
            }
        )
    return out


def _gather(
    db: DatabaseManager,
    user: dict,
    q: str = "",
    only_sections: Optional[List[str]] = None,
    sort_field: str = "",
    order: str = "asc",
    date_field: str = "",
) -> List[Dict[str, Any]]:
    uid = None if is_admin(user) else int(user["id"])
    admin = is_admin(user)
    want = set(only_sections or [])
    gathered: List[Dict[str, Any]] = []
    for table in DatabaseManager.JSON_TABLES:
        if want and table not in want:
            continue
        try:
            rows, _ = db.query_json_records(
                table, owner_id=uid, is_admin=admin, q=q, page_size=MAX_PER_SECTION
            )
        except Exception:
            continue
        for r in rows:
            dj = r.get("data_json") or {}
            gathered.append(
                {
                    "section": table,
                    "section_label": _label(table),
                    "id": r["id"],
                    "data": dj,
                    "created_at": r.get("created_at"),
                    "user_id": r.get("user_id", 0),
                }
            )
    for t in db.get_custom_tables(uid, admin):
        if want and t["key"] not in want:
            continue
        try:
            rows, _ = db.query_custom_records(
                t["key"], owner_id=uid, is_admin=admin, q=q, page_size=MAX_PER_SECTION
            )
        except Exception:
            continue
        for r in rows:
            dj = r.get("data_json") or {}
            gathered.append(
                {
                    "section": t["key"],
                    "section_label": t["label"],
                    "id": r["id"],
                    "data": dj,
                    "created_at": r.get("created_at"),
                    "user_id": r.get("user_id", 0),
                }
            )
    return _sort_rows(gathered, sort_field, order, date_field)


def _sort_rows(rows, sort_field: str, order: str, date_field: str = ""):
    direction = -1 if str(order).lower() == "desc" else 1

    def keyfn(r: Dict[str, Any]):
        if sort_field == "id":
            v = r["id"]
        elif sort_field == "section":
            v = r["section_label"]
        elif sort_field == "created_at":
            v = r.get("created_at") or r.get("data", {}).get("Дата", "")
        elif sort_field:
            raw = r["data"].get(sort_field, r["data"].get(date_field, ""))
            if isinstance(raw, (int, float)):
                v = raw
            else:
                v = str(raw or "")
        else:
            v = r["section_label"]
        return v

    try:
        rows.sort(key=keyfn, reverse=direction < 0)
    except Exception:
        rows.sort(key=lambda r: str(r["section_label"]))
    return rows


def _union_columns(
    db: DatabaseManager, user: dict, only_sections: Optional[List[str]] = None
) -> List[str]:
    """Объединение имён колонок всех доступных разделов (для экспорта)."""
    names: List[str] = []
    seen = set()
    for sec in _sections(db, user):
        if only_sections and sec["section"] not in only_sections:
            continue
        for n in sec["columns"]:
            if n not in seen:
                seen.add(n)
                names.append(n)
    return names


class ExportIn(BaseModel):
    q: str = ""
    sections: List[str] = []
    columns: List[str] = []
    format: str = "xlsx"  # xlsx | csv | json
    title: str = "Реестр Всё"


@router.get("/sections")
def union_sections(db=Depends(get_db), user=Depends(get_current_user)):
    return {"sections": _sections(db, user)}


@router.get("/records")
def union_records(
    request: Request,
    page: int = 1,
    page_size: int = 50,
    q: str = "",
    sort_by: str = "",
    order: str = "asc",
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    filters = [
        v for k, v in request.query_params.multi_items() if k == "f_section" and v
    ]
    only = [x for x in filter(None, (i for v in filters for i in v.split(",")))]
    order_cast = request.query_params.get("order_cast", "")
    gathered = _gather(
        db,
        user,
        q=q,
        only_sections=only,
        sort_field=sort_by,
        order=order,
        date_field=request.query_params.get("date_field", ""),
    )
    total = len(gathered)
    page = max(1, int(page))
    page_size = min(500, max(1, int(page_size)))
    start = (page - 1) * page_size
    items = gathered[start : start + page_size]
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "columns": _union_columns(db, user, only_sections=only),
        "sections": _sections(db, user),
        "_t": datetime.now().isoformat(),
    }


@router.post("/export")
def union_export(body: ExportIn, db=Depends(get_db), user=Depends(get_current_user)):
    section_map = {s["section"]: s for s in _sections(db, user)}
    gathered = _gather(db, user, q=body.q, only_sections=body.sections or None)
    cols = [c for c in body.columns] or _union_columns(
        db, user, only_sections=body.sections or None
    )
    rows = []
    for it in gathered:
        row = {"Раздел": it["section_label"]}
        for name in cols:
            v = it["data"].get(name, "")
            row[name] = v
        rows.append(row)
    header = ["Раздел"] + [c for c in cols if c != "Раздел"]
    fmt = (body.format or "xlsx").lower()
    if fmt == "json":
        data = json.dumps(
            {"title": body.title, "columns": header, "rows": rows},
            ensure_ascii=False,
            indent=2,
        )
        return _respond(
            data.encode("utf-8"),
            f"union_{datetime.now():%Y%m%d}.json",
            "application/json",
        )
    if fmt == "csv":
        import csv as _csv

        buf = io.StringIO()
        w = _csv.writer(buf, delimiter=";", quoting=_csv.QUOTE_MINIMAL)
        w.writerow(header)
        for row in rows:
            w.writerow([str(row.get(c, "") or "") for c in header])
        return _respond(
            buf.getvalue().encode("utf-8-sig"),
            f"union_{datetime.now():%Y%m%d}.csv",
            "text/csv",
        )
    data = _xlsx_simple(body.title, header, rows)
    return _respond(
        data,
        f"union_{datetime.now():%Y%m%d}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _xlsx_simple(title: str, header: List[str], rows: List[Dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or "Реестр")[:31]
    head_fill = PatternFill("solid", fgColor="6366F1")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    for i, c in enumerate(header, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for ri, row in enumerate(rows, 2):
        for ci, c in enumerate(header, 1):
            v = row.get(c, "")
            if isinstance(v, (dict, list)):
                v = JsonUtils.dumps(v)
            ws.cell(row=ri, column=ci, value=str(v) if v is not None else "")
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = max(
            14, min(40, len(header[i - 1]) + 6)
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _respond(data: bytes, filename: str, mime: str):
    return StreamingResponse(
        io.BytesIO(data),
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

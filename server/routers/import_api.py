"""Импорт-мастер: файл/буфер → маппинг → сухой прогон → выполнение →
отмена, шаблоны, фото из ZIP, история."""

import csv
import io
import json
import os
import uuid
import zipfile
from datetime import datetime, date
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from server.deps import get_db, get_current_user, is_admin
from services.database import DatabaseManager
from app_core.utils import JsonUtils

router = APIRouter(prefix="/api/import", tags=["import"])

STASH: Dict[str, Dict[str, Any]] = {}
STASH_TTL = 3600 * 2
MAX_ROWS = 5000


# ── разбор файлов ──


def _norm_date(v: Any) -> Optional[str]:
    """Нормализация дат: Excel serial/datetime/строки → ДД.ММ.ГГГГ.
    None — если разобрать не удалось (нужна для отчёта ошибок)."""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    s = str(v or "").strip()
    if not s:
        return ""
    fmts = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y.%m.%d", "%d.%m.%y")
    for f in fmts:
        try:
            return datetime.strptime(s, f).strftime("%d.%m.%Y")
        except ValueError:
            continue
    try:
        num = float(s)
        if 20000 < num < 60000:
            base = datetime(1899, 12, 30)
            return (base + __import__("datetime").timedelta(days=int(num))).strftime(
                "%d.%m.%Y"
            )
    except ValueError:
        pass
    return None


def _stash_put(content: bytes, filename: str) -> str:
    fid = uuid.uuid4().hex[:16]
    STASH[fid] = {"content": content, "filename": filename, "ts": datetime.now()}
    # чистка старых
    for k in list(STASH):
        if (datetime.now() - STASH[k]["ts"]).total_seconds() > STASH_TTL:
            STASH.pop(k, None)
    return fid


def _parse_sheet(
    content: bytes, filename: str, sheet: str = ""
) -> Tuple[List[str], List[List[Any]]]:
    fn = filename.lower()
    if fn.endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            vals = ["" if v is None else v for v in r]
            if any(str(v).strip() for v in vals):
                rows.append(vals)
        wb.close()
        headers = [str(h).strip() for h in (rows[0] if rows else [])]
        return headers, rows[1:] if rows else []
    # CSV / TSV (буфер обмена)
    text = content.decode("utf-8-sig", errors="replace")
    sep = (
        "\t"
        if "\t" in text.splitlines()[0]
        else (";" if ";" in text.splitlines()[0] else ",")
    )
    reader = csv.reader(text.splitlines(), delimiter=sep)
    rows = [r for r in reader if any(str(x).strip() for x in r)]
    headers = [str(h).strip() for h in (rows[0] if rows else [])]
    return headers, rows[1:] if rows else []


def _sheet_names(content: bytes, filename: str) -> List[str]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        names = sorted(wb.sheetnames)
        wb.close()
        return names
    return [""]


# ── эндпоинты ──


class UploadTextIn(BaseModel):
    text: str
    filename: str = "clipboard.csv"


@router.post("/upload")
async def upload_file(
    file: Optional[UploadFile] = None,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    if file is None:
        raise HTTPException(400, "Файл не передан")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "Файл больше 10 МБ")
    fn = file.filename or "import.csv"
    if not fn.lower().endswith((".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".zip")):
        raise HTTPException(400, "Поддерживаются .xlsx и .csv")
    fid = _stash_put(content, fn)
    return {"file_id": fid, "filename": fn, "sheets": _sheet_names(content, fn)}


@router.post("/upload_text")
def upload_text(body: UploadTextIn, db=Depends(get_db), user=Depends(get_current_user)):
    fid = _stash_put(body.text.encode("utf-8"), body.filename)
    return {
        "file_id": fid,
        "filename": body.filename,
        "sheets": _sheet_names(body.text.encode("utf-8"), body.filename),
    }


@router.get("/preview/{fid}")
def preview(
    fid: str,
    sheet: str = "",
    limit: int = 5,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    st = STASH.get(fid)
    if not st:
        raise HTTPException(404, "Файл истёк, загрузите заново")
    headers, rows = _parse_sheet(st["content"], st["filename"], sheet)
    return {"headers": headers, "total": len(rows), "rows": rows[:limit]}


class MergeRule(BaseModel):
    target: str
    parts: List[str]
    sep: str = " "


class AnalyzeIn(BaseModel):
    file_id: str
    sheet: str = ""
    table: str
    mapping: Dict[str, str] = {}  # колонка файла → колонка таблицы
    merges: List[MergeRule] = []
    key_field: str = ""  # для обновления существующих
    mode: str = "upsert"  # upsert | insert
    limit: int = 5000


def _build_records(
    body: AnalyzeIn,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[str]]:
    st = STASH.get(body.file_id)
    if not st:
        raise HTTPException(404, "Файл истёк, загрузите заново")
    headers, rows = _parse_sheet(st["content"], st["filename"], body.sheet)
    db = DatabaseManager()
    sys_cols = {
        c["name"]: c["type"]
        for c in db.get_columns_config(body.table)
        if c["name"] != "ID"
    }
    records: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    for i, raw in enumerate(rows):
        rec: Dict[str, Any] = {}
        row_error = None
        for ci, h in enumerate(headers):
            target = body.mapping.get(h)
            if not target or target not in sys_cols:
                continue
            v = raw[ci] if ci < len(raw) else ""
            if sys_cols[target] in ("Дата", "Годен до", "Дата проведения"):
                v = _norm_date(v)
                if v is None:
                    row_error = (
                        f"строка {i + 2}: «{v if v else ''}» — "
                        f"неверная дата для «{target}» "
                        f"(нужно ДД.ММ.ГГГГ)"
                    )
                    break
            elif sys_cols[target] == "Число":
                sv = str(v).strip()
                if sv:
                    try:
                        v = float(sv.replace(",", "."))
                    except ValueError:
                        row_error = f"строка {i + 2}: «{v}» не число для «{target}»"
                        break
            rec[target] = v
        if row_error:
            errors.append({"row": i + 2, "error": row_error})
            continue
        # объединение колонок
        for m in body.merges:
            if m.target not in sys_cols:
                continue
            parts = []
            for p in m.parts:
                ci = headers.index(p) if p in headers else -1
                if ci >= 0:
                    parts.append(str(raw[ci]) if ci < len(raw) else "")
            rec[m.target] = m.sep.join(x for x in parts if x.strip())
        if not any(str(v).strip() for v in rec.values()):
            continue
        records.append(rec)
    return records, errors, headers


@router.post("/analyze")
def analyze(body: AnalyzeIn, db=Depends(get_db), user=Depends(get_current_user)):
    _table_ok(body.table)
    records, errors, headers = _build_records(body)
    uid = int(user["id"])
    admin = is_admin(user)
    to_update = 0
    key_idx: Dict[str, Dict[str, Any]] = {}
    if body.mode == "upsert" and body.key_field:
        kf = body.key_field
        existing = db.query_json_records(
            body.table,
            owner_id=None if admin else uid,
            is_admin=admin,
            page_size=MAX_ROWS,
        )[0]
        for ex in existing:
            kv = str(ex.get(kf, "")).strip().lower()
            if kv:
                key_idx[kv] = ex
        for r in records:
            kv = str(r.get(kf, "")).strip().lower()
            if kv and kv in key_idx:
                to_update += 1
    to_create = len(records) - to_update
    return {
        "total": len(records),
        "to_create": to_create,
        "to_update": to_update,
        "errors": errors[:50],
        "error_count": len(errors),
        "sample": records[:3],
    }


class RunIn(AnalyzeIn):
    dry: bool = False


@router.post("/run")
def run_import(body: RunIn, db=Depends(get_db), user=Depends(get_current_user)):
    _table_ok(body.table)
    records, errors, headers = _build_records(body)
    uid = int(user["id"])
    admin = is_admin(user)
    created_ids: List[int] = []
    updated_snaps: List[Dict[str, Any]] = []
    updated_n = 0
    created_n = 0

    key_idx: Dict[str, Dict[str, Any]] = {}
    if body.mode == "upsert" and body.key_field:
        kf = body.key_field
        existing = db.query_json_records(
            body.table,
            owner_id=None if admin else uid,
            is_admin=admin,
            page_size=MAX_ROWS,
        )[0]
        for ex in existing:
            kv = str(ex.get(kf, "")).strip().lower()
            if kv:
                key_idx[kv] = ex

    for r in records:
        kv = str(r.get(body.key_field, "")).strip().lower() if body.key_field else ""
        target = key_idx.get(kv) if (body.mode == "upsert" and kv) else None
        if target:
            prev = db.get_json_record(body.table, target["id"])
            prev_data = dict(prev.get("data_json") or {}) if prev else {}
            merged = {**prev_data, **r}
            db.save_json_record(body.table, target["id"], merged, user_id=uid)
            updated_snaps.append({"id": target["id"], "prev": prev_data})
            updated_n += 1
        else:
            nid = db.save_json_record(body.table, 0, r, user_id=uid)
            created_ids.append(nid)
            created_n += 1

    details = {
        "created_ids": created_ids,
        "updated": updated_snaps,
        "error_count": len(errors),
    }
    cur = db.execute(
        "INSERT INTO import_history (table_name, source_file, imported, "
        "updated, errors, details) VALUES (?, ?, ?, ?, ?, ?)",
        (
            body.table,
            STASH.get(body.file_id, {}).get("filename", ""),
            created_n,
            updated_n,
            len(errors),
            JsonUtils.dumps(details),
        ),
    )
    db.commit()
    import_id = int(cur.lastrowid)
    return {
        "import_id": import_id,
        "created": created_n,
        "updated": updated_n,
        "errors": errors[:50],
        "error_count": len(errors),
    }


class UndoIn(BaseModel):
    import_id: int


@router.post("/undo")
def undo_import(body: UndoIn, db=Depends(get_db), user=Depends(get_current_user)):
    h = db.fetch_one("SELECT * FROM import_history WHERE id=?", (body.import_id,))
    if not h:
        raise HTTPException(404, "Импорт не найден")
    try:
        details = JsonUtils.loads(h["details"] or "{}")
    except Exception:
        details = {}
    undone = 0
    for cid in details.get("created_ids", []):
        if db.fetch_one(f"SELECT id FROM {h['table_name']} WHERE id=?", (cid,)):
            db.delete_json_record(h["table_name"], cid)
            undone += 1
    for snap in details.get("updated", []):
        db.save_json_record(
            h["table_name"], snap["id"], snap.get("prev", {}), user_id=int(user["id"])
        )
        undone += 1
    db.execute("DELETE FROM import_history WHERE id=?", (body.import_id,))
    db.commit()
    return {"undone": undone}


@router.get("/history")
def history(limit: int = 20, db=Depends(get_db), user=Depends(get_current_user)):
    rows = db.fetch_all(
        "SELECT id, timestamp, table_name, source_file, imported, updated, "
        "errors FROM import_history ORDER BY id DESC LIMIT ?",
        (limit,),
    )
    return {"items": [dict(r) for r in rows]}


@router.get("/template/{table}")
def template_xlsx(table: str, db=Depends(get_db), user=Depends(get_current_user)):
    if table not in DatabaseManager.JSON_TABLES:
        raise HTTPException(404, "Неизвестная таблица")
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table
    cols = [
        c
        for c in db.get_columns_config(table)
        if c["name"] != "ID" and c["type"] not in ("Медиа", "Фото")
    ]
    for i, c in enumerate(cols, 1):
        ws.cell(row=1, column=i, value=c["name"])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=template_{table}.xlsx"},
    )


# ── Фото из ZIP ──


class PhotosIn(BaseModel):
    file_id: str
    table: str
    match_by: str = "ФИО"


@router.post("/photos_zip")
def photos_zip(body: PhotosIn, db=Depends(get_db), user=Depends(get_current_user)):
    if body.table not in DatabaseManager.JSON_TABLES:
        raise HTTPException(404, "Неизвестная таблица")
    st = STASH.get(body.file_id)
    if not st:
        raise HTTPException(404, "ZIP истёк, загрузите заново")
    if not st["filename"].lower().endswith(".zip"):
        raise HTTPException(400, "Нужен .zip архив")
    from app_core.config import RUNTIME_PATHS

    media_dir = str(RUNTIME_PATHS.media_dir)
    os.makedirs(media_dir, exist_ok=True)
    uid = int(user["id"])
    admin = is_admin(user)
    rows = db.query_json_records(
        body.table, None if admin else uid, is_admin=admin, page_size=MAX_ROWS
    )[0]
    ext_ok = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
    matched, unmatched = 0, []
    with zipfile.ZipFile(io.BytesIO(st["content"])) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            base = os.path.basename(info.filename)
            stem, ext = os.path.splitext(base)
            if ext.lower() not in ext_ok:
                continue
            stem_clean = stem.strip()
            target = next(
                (
                    r
                    for r in rows
                    if str(r.get(body.match_by, "")).strip() == stem_clean
                ),
                None,
            )
            if not target:
                unmatched.append(base)
                continue
            name = f"{uuid.uuid4().hex[:12]}{ext.lower()}"
            with open(os.path.join(media_dir, name), "wb") as fh:
                fh.write(zf.read(info))
            data = dict(target.get("data_json") or {})
            data["Фото"] = f"/media/{name}"
            db.save_json_record(body.table, target["id"], data, user_id=uid)
            matched += 1
    return {
        "matched": matched,
        "unmatched": unmatched[:30],
        "unmatched_count": len(unmatched),
    }


def _table_ok(table: str) -> None:
    if table not in DatabaseManager.JSON_TABLES:
        raise HTTPException(404, f"Неизвестная таблица: {table}")

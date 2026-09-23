"""Экспорт: xlsx (с фото/стилями), csv, json, docx, ZIP-бандл всех таблиц."""

import io
from datetime import datetime
import json
import os
import uuid
import zipfile
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from server.deps import get_db, get_current_user, is_admin
from services.database import DatabaseManager
from app_core.utils import JsonUtils

router = APIRouter(prefix="/api/export", tags=["export"])


def _merged_columns(
    db: DatabaseManager, table: str, user: dict
) -> List[Dict[str, Any]]:
    """Системные колонки + личные колонки пользователя."""
    cols = [
        {**c, "source": "sys"}
        for c in db.get_columns_config(table)
        if c["name"] != "ID"
    ]
    if not table.startswith("u_"):
        rows = db.fetch_all(
            "SELECT name, type, template, visible FROM user_columns "
            "WHERE table_key=? AND user_id=? ORDER BY position, id",
            (table, int(user["id"])),
        )
        cols += [
            {
                "name": r["name"],
                "type": r["type"],
                "template": r["template"] or "",
                "visible": r["visible"],
                "source": "my",
            }
            for r in rows
        ]
    return cols


def _collect_rows(
    db: DatabaseManager,
    table: str,
    user: dict,
    ids: Optional[List[int]],
    request: Optional[Request],
    columns: List[str],
) -> List[Dict[str, Any]]:
    """Строки по диапазону: ids → фильтры запроса → всё."""
    uid = None if is_admin(user) else int(user["id"])
    admin = is_admin(user)
    if ids:
        out = []
        for rid in ids:
            # IDOR-фикс (аудит 7.2 п.1): экспорт по ids — только свои записи.
            if not db.user_can_access(table, int(rid), int(user["id"]), admin):
                continue
            rec = (
                db.get_json_record(table, int(rid))
                if table in DatabaseManager.JSON_TABLES
                else db.get_custom_record(table, int(rid))
            )
            if rec:
                out.append(rec)
        return out
    if request is not None:
        q = request.query_params.get("q", "")
        filters: Dict[str, List[str]] = {}
        for k, v in request.query_params.multi_items():
            if k.startswith("f_") and v:
                filters[k[2:]] = [x for x in v.split(",") if x]
        sort_by = request.query_params.get("sort_by", "")
        order = request.query_params.get("order", "asc")
        if table in DatabaseManager.JSON_TABLES:
            rows, _ = db.query_json_records(
                table,
                owner_id=uid,
                is_admin=admin,
                q=q,
                filters=filters,
                sort_by=sort_by,
                order=order,
                page_size=MAXX,
            )
        else:
            rows, _ = db.query_custom_records(
                table,
                owner_id=uid,
                is_admin=admin,
                q=q,
                filters=filters,
                sort_by=sort_by,
                order=order,
                page_size=MAXX,
            )
        return rows
    if table in DatabaseManager.JSON_TABLES:
        return db.query_json_records(
            table, owner_id=uid, is_admin=admin, page_size=MAXX
        )[0]
    return db.query_custom_records(table, owner_id=uid, is_admin=admin, page_size=MAXX)[
        0
    ]


MAXX = 20000


class ExportIn(BaseModel):
    table: str
    columns: List[str] = []  # пусто → все доступные
    ids: List[int] = []  # выбранные строки
    with_photos: bool = False
    title: str = ""


def _resolve(
    db: DatabaseManager,
    table: str,
    user: dict,
    body: ExportIn,
    request: Optional[Request],
):
    if table.startswith("u_"):
        t = db.get_custom_table(table)
        if not t:
            raise HTTPException(404, "Таблица не найдена")
        owner = int(t["user_id"] or 0)
        if owner != 0 and owner != int(user["id"]) and not is_admin(user):
            raise HTTPException(403, "Нет доступа")
        try:
            sys_cols = JsonUtils.loads(t["columns_json"] or "[]")
        except Exception:
            sys_cols = []
        sys_cols = [
            {"name": c["name"], "type": c.get("type", "Текст")}
            for c in sys_cols
            if c.get("name") != "ID"
        ]
        merged = [{**c, "source": "tbl"} for c in sys_cols]
    else:
        if table not in DatabaseManager.JSON_TABLES:
            raise HTTPException(404, "Неизвестная таблица")
        merged = _merged_columns(db, table, user)
    if body.columns:
        want = set(body.columns)
        merged = [c for c in merged if c["name"] in want]
    rows = _collect_rows(
        db, table, user, body.ids or None, request, [c["name"] for c in merged]
    )
    return merged, rows


def _xlsx(
    merged, rows, table: str, title: str, with_photos: bool, db: DatabaseManager
) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or table)[:31]
    head_fill = PatternFill("solid", fgColor="6366F1")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Border(bottom=Side(style="thin", color="E5E7EB"))
    names = [c["name"] for c in merged]
    for i, c in enumerate(merged, 1):
        cell = ws.cell(row=1, column=i, value=c["name"])
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(
            14, min(40, len(c["name"]) + 6)
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(merged))}1"

    photo_cells = []  # (row, col, media_path)
    for ri, row in enumerate(rows, 2):
        for ci, c in enumerate(merged, 1):
            v = row.get(c["name"], "")
            if c.get("type") == "Вычисляемая" and c.get("template"):
                v = str(c["template"]).replace("{", "{").format()
                tmpl = c["template"]
                import re as _re

                v = _re.sub(
                    r"\{([^}]+)\}", lambda m: str(row.get(m.group(1), "") or ""), tmpl
                ).strip()
            if isinstance(v, (dict, list)):
                v = JsonUtils.dumps(v)
            cell = ws.cell(row=ri, column=ci, value=str(v) if v is not None else "")
            cell.alignment = Alignment(wrap_text=False, vertical="center")
            cell.border = thin
        if with_photos:
            photo = str(row.get("Фото", "") or "")
            if photo.startswith("/media/"):
                photo_cells.append((ri, names.index("Фото") + 1, photo))
    # автоширина по содержимому
    for ci, c in enumerate(merged, 1):
        maxlen = len(c["name"])
        for row in rows[:300]:
            v = str(row.get(c["name"], "") or "")
            maxlen = max(maxlen, min(len(v), 60))
        ws.column_dimensions[get_column_letter(ci)].width = max(14, min(46, maxlen + 4))

    if with_photos and photo_cells:
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image as XLImage

        tmpdir = tempfile.gettempdir()
        for ri, ci, path in photo_cells[:100]:
            fs_path = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                path.lstrip("/"),
            )
            if not os.path.isfile(fs_path):
                continue
            try:
                img = PILImage.open(fs_path)
                img.thumbnail((64, 64))
                tmp = os.path.join(tmpdir, f"exp_{uuid.uuid4().hex[:8]}.png")
                img.save(tmp, "PNG")
                xl = XLImage(tmp)
                xl.width, xl.height = 48, 48
                ws.add_image(xl, f"{get_column_letter(ci)}{ri}")
            except Exception:
                pass
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(merged, rows, table: str, title: str) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(title or table, level=1)
    t = doc.add_table(rows=1, cols=len(merged))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(merged):
        t.rows[0].cells[i].text = c["name"]
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(merged):
            v = row.get(c["name"], "")
            cells[i].text = str(v) if v is not None else ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class ExportIn(BaseModel):
    table: str
    columns: List[str] = []
    ids: List[int] = []
    with_photos: bool = False
    title: str = ""
    zip_photos: bool = False


def _response(data: bytes, filename: str, mime: str):
    return StreamingResponse(
        io.BytesIO(data),
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/xlsx/{table}")
def export_xlsx(
    table: str,
    body: ExportIn,
    request: Request,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    merged, rows = _resolve(db, table, user, body, request)
    data = _xlsx(merged, rows, table, body.title or table, body.with_photos, db)
    return _response(
        data,
        f"{table}_{datetime.now():%Y%m%d}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.post("/csv/{table}")
def export_csv(
    table: str,
    body: ExportIn,
    request: Request,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    merged, rows = _resolve(db, table, user, body, request)
    import csv as _csv

    buf = io.StringIO()
    w = _csv.writer(buf, delimiter=";", quoting=_csv.QUOTE_MINIMAL)
    w.writerow([c["name"] for c in merged])
    for row in rows:
        w.writerow([str(row.get(c["name"], "") or "") for c in merged])
    return _response(
        buf.getvalue().encode("utf-8-sig"),
        f"{table}_{datetime.now():%Y%m%d}.csv",
        "text/csv",
    )


@router.post("/json/{table}")
def export_json(
    table: str,
    body: ExportIn,
    request: Request,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    merged, rows = _resolve(db, table, user, body, request)
    names = [c["name"] for c in merged]
    items = [{k: row.get(k) for k in names if k in row} for row in rows]
    data = json.dumps({"table": table, "items": items}, ensure_ascii=False, indent=2)
    return _response(
        data.encode("utf-8"),
        f"{table}_{datetime.now():%Y%m%d}.json",
        "application/json",
    )


@router.post("/docx/{table}")
def export_docx(
    table: str,
    body: ExportIn,
    request: Request,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    merged, rows = _resolve(db, table, user, body, request)
    data = _docx(merged, rows, table, body.title or table)
    return _response(
        data,
        f"{table}_{datetime.now():%Y%m%d}.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@router.post("/xlsx_photos_zip/{table}")
def export_photos_zip(
    table: str,
    body: ExportIn,
    request: Request,
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    """xlsx с миниатюрами + ZIP всех фото."""
    merged, rows = _resolve(db, table, user, body, request)
    xlsx_data = _xlsx(merged, rows, table, body.title or table, True, db)
    from app_core.config import RUNTIME_PATHS

    media_dir = str(RUNTIME_PATHS.media_dir)
    zbuf = io.BytesIO()
    n = 0
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{table}.xlsx", xlsx_data)
        for row in rows:
            photo = str(row.get("Фото", "") or "")
            if not photo.startswith("/media/"):
                continue
            fs = os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                photo.lstrip("/"),
            )
            if os.path.isfile(fs):
                stem = str(
                    row.get("ФИО", row.get("Наименование", row.get("id", "photo")))
                )
                zf.writestr(
                    f"{stem}{os.path.splitext(photo)[1]}", open(fs, "rb").read()
                )
                n += 1
    return _response(zbuf.getvalue(), f"{table}_with_photos.zip", "application/zip")


@router.post("/bundle")
def export_bundle(db=Depends(get_db), user=Depends(get_current_user)):
    """ZIP с xlsx всех системных таблиц (личные данные пользователя)."""
    uid = None if is_admin(user) else int(user["id"])
    admin = is_admin(user)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
        for table in sorted(DatabaseManager.JSON_TABLES):
            try:
                merged = _merged_columns(db, table, user)
                rows = db.query_json_records(
                    table, owner_id=uid, is_admin=admin, page_size=MAXX
                )[0]
                data = _xlsx(merged, rows, table, table, False, db)
                zf.writestr(f"{table}.xlsx", data)
            except Exception:
                continue
    return _response(
        zbuf.getvalue(), f"suot_export_{datetime.now():%Y%m%d}.zip", "application/zip"
    )
